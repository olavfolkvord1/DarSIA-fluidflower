from __future__ import annotations

import os
import time
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
pytest.importorskip("PySide6")

import numpy as np
from darsia.fluidflower.data_integrity import IntegrityReportStore, IntegrityScanner
from darsia.fluidflower.geometry_facies import (
    GeometryRevisionStore,
    geometry_config_digest,
    load_geometry_preview,
    validate_geometry_config,
)
from darsia.fluidflower.gui import (
    DIRECTORY_DIALOG_OPTIONS,
    AddRunsDialog,
    CalibrationPage,
    DataIntegrityPage,
    EstablishedSourceSetupEditor,
    ExecutionPage,
    FluidFlowerMainWindow,
    GeometryFaciesPage,
    HomePage,
    ImageCorrectionPage,
    ImageSourcesPage,
    MassAnalysisPage,
    NewSeriesWizard,
    OpticalCalibrationPage,
    PhaseInterpretationPage,
    PhaseModelStudioWindow,
    PostAnalysisPage,
    ProjectWorkspace,
    RigProtocolPage,
    SeriesIdentityPage,
    TimelinePage,
    UpdateRunSourcesDialog,
    WassersteinPage,
    application_icon,
)
from darsia.fluidflower.modules import MODULE_DEFINITIONS
from darsia.fluidflower.page_registry import (
    FluidFlowerPagePlugin,
    FluidFlowerSubpagePlugin,
    default_page_registry,
    default_subpage_registry,
)
from darsia.fluidflower.pages.image_correction.page import ZoomableImageView
from darsia.fluidflower.pages.setup.global_run_structure import (
    GlobalRunStructureDialog,
    _template_editor_values,
)
from darsia.fluidflower.pages.setup.incremental_runs import (
    _cross_run_mapping_issues,
)
from darsia.fluidflower.pages.setup.source_layout_dialog import (
    RunLayoutDialog,
    SeriesRunsDialog,
)
from darsia.fluidflower.project import FluidFlowerProject, ProjectRegistry
from darsia.fluidflower.protocol_schema import write_protocol_workbook_template
from darsia.fluidflower.setup_drafts import SetupDraftStore
from darsia.fluidflower.source_discovery import ROLE_INJECTION, ROLE_PROTOCOL_TABLE
from darsia.fluidflower.timeline import TimelineStore
from darsia.fluidflower.ui.interaction_safety import (
    ComboBoxWheelGuard,
    install_interaction_safety,
)
from PIL import Image
from PySide6.QtCore import QEvent, Qt
from PySide6.QtWidgets import (
    QApplication,
    QComboBox,
    QFileDialog,
    QLabel,
    QMessageBox,
    QScrollArea,
    QSpinBox,
    QTableWidgetSelectionRange,
)


@pytest.fixture(scope="module")
def qt_app() -> QApplication:
    app = QApplication.instance() or QApplication([])
    return app


def test_series_identity_page_has_valid_editable_default(
    qt_app: QApplication,
) -> None:
    page = SeriesIdentityPage()
    assert page.template_id() == "ff_um"
    assert page.name.text() == "FF_UM"
    assert page.series_id.text() == "ff_um"
    assert page.target_path().name == "ff-um"
    assert page.default_result_path() == page.target_path() / "results"
    assert "raw experiment data" in page.storage_explanation.text().lower()
    assert page.isComplete()


def test_template_change_updates_unedited_defaults(qt_app: QApplication) -> None:
    page = SeriesIdentityPage()
    page.template.setCurrentIndex(page.template.findData("ff_ac"))
    assert page.name.text() == "FF_AC"
    assert page.series_id.text() == "ff_ac"


def test_protocol_example_controls_follow_selected_organization(
    qt_app: QApplication,
) -> None:
    identity = SeriesIdentityPage()
    page = RigProtocolPage(identity)

    page.protocol_layout.setCurrentIndex(page.protocol_layout.findData("combined"))
    assert not page.protocol.directory
    assert "combined-workbook" in page.create_protocol_template.text()

    page.protocol_layout.setCurrentIndex(
        page.protocol_layout.findData("per_run_sheets")
    )
    assert not page.protocol.directory
    assert "per-run-sheet" in page.create_protocol_template.text()

    page.protocol_layout.setCurrentIndex(page.protocol_layout.findData("per_run_files"))
    assert not page.protocol.directory
    assert page.protocol.isHidden()
    assert not page.protocol_discovery_note.isHidden()
    assert "per-run-file" in page.create_protocol_template.text()
    assert "inside each experiment folder" in page.protocol_layout.currentText()
    page.close()


def test_apparatus_menu_shows_nominal_medium_and_large_dimensions(
    qt_app: QApplication,
) -> None:
    page = RigProtocolPage(SeriesIdentityPage())

    assert "0.90 x 0.49 m" in page.rig_type.itemText(page.rig_type.findData("medium"))
    assert "2.80 x 1.50 m" in page.rig_type.itemText(page.rig_type.findData("large"))
    page.close()


def test_setup_wizard_merges_rig_protocol_and_image_sources(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    wizard = NewSeriesWizard(draft_store=SetupDraftStore(tmp_path / "drafts"))

    assert len(wizard.pageIds()) == 3
    assert wizard.pageIds() == [0, 1, 2]
    assert wizard.rig_protocol is wizard.images.rig_protocol
    assert wizard.images.title() == "Apparatus, protocol and image sources"
    assert wizard.rig_protocol.protocol_layout.currentData() == "per_run_files"
    wizard.close()


def test_combined_setup_page_scrolls_instead_of_compressing_controls(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    wizard = NewSeriesWizard(draft_store=SetupDraftStore(tmp_path / "drafts"))
    wizard.resize(980, 700)
    wizard.show()
    qt_app.processEvents()
    wizard.next()
    qt_app.processEvents()

    assert wizard.images.scroll_area.verticalScrollBar().maximum() > 0
    assert wizard.images.scroll_area.horizontalScrollBar().maximum() > 0
    wizard.close()


def test_main_window_starts_without_recent_projects(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    window = FluidFlowerMainWindow(registry=ProjectRegistry(tmp_path / "registry.json"))
    assert window.windowTitle() == "DarSIA · FluidFlower Experiments"
    assert window.home.table.rowCount() == 0
    assert not hasattr(window.home, "shared_model_training")
    window.close()
    window.deleteLater()
    qt_app.processEvents()


def test_phase_model_training_runs_in_standalone_studio(
    qt_app: QApplication,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "darsia.fluidflower.ui.model_training.discover_phase_review_sessions",
        list,
    )
    monkeypatch.setattr(
        "darsia.fluidflower.ui.model_training.discover_phase_model_history",
        lambda _root: [],
    )
    window = PhaseModelStudioWindow()
    page = window.training_page

    assert window.windowTitle() == "DarSIA · Phase Model Studio"
    assert not page.start_button.isEnabled()
    assert page.pixels.maximum() == 1_000_000
    assert page.tabs.count() == 3
    assert page.tabs.tabText(0) == "Review and label evidence"
    assert page.tabs.tabText(1) == "Train a new model version"
    assert page.tabs.tabText(2) == "Model history and validation"
    assert page.open_blind_labeling.text() == "Start new blind phase-labeling session"
    assert page.open_model_correction.text() == "Start new model-validation session"
    window.close()
    window.deleteLater()
    qt_app.processEvents()


def test_phase_model_studio_exposes_resumable_review_modes(
    qt_app: QApplication,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    common = {
        "campaign_id": "shared-v5",
        "reviewer": "Olav Folkvord",
        "ffac_root": "C:/campaign/ffac",
        "ffum_root": "C:/campaign/ffum",
        "combined_root": "C:/campaign/combined",
        "local_review_root": "C:/reviews",
        "modified_timestamp": 1.0,
        "reviewed_experiments": 2,
        "frame_verdicts": 0,
    }
    sessions = [
        {
            **common,
            "workspace": "current_comparison",
            "workspace_title": "Open model correction",
            "session_path": "C:/reviews/open/review_session.json",
            "decision_path": "C:/reviews/open/decisions.json",
            "manual_regions": 57,
            "blind_region_labels": 0,
            "editing_locked": True,
        },
        {
            **common,
            "workspace": "training",
            "workspace_title": "Blind phase labeling",
            "session_path": "C:/reviews/blind/review_session.json",
            "decision_path": "C:/reviews/blind/decisions.json",
            "manual_regions": 0,
            "blind_region_labels": 31,
        },
    ]
    monkeypatch.setattr(
        "darsia.fluidflower.ui.model_training.discover_phase_review_sessions",
        lambda: sessions,
    )
    monkeypatch.setattr(
        "darsia.fluidflower.ui.model_training.discover_phase_model_history",
        lambda _root: [],
    )

    window = PhaseModelStudioWindow()
    page = window.training_page

    assert not page.open_blind_labeling.isEnabled()
    assert page.open_blind_labeling.text() == "Start new blind phase-labeling session"
    assert not page.open_model_correction.isEnabled()
    assert page.open_model_correction.text() == "Start new model-validation session"
    page.ffac_campaign.setText("C:/selected-model/ffac")
    page.ffum_campaign.setText("C:/selected-model/ffum")
    page.combined_campaign.setText("C:/selected-model/combined")
    page.session_scope.setCurrentIndex(1)
    assert page.resume_session_button.isEnabled()
    assert page.resume_session_button.text() == "Review locked evidence (read-only)"
    assert page.ffac_campaign.text() == "C:/selected-model/ffac"
    assert page.ffum_campaign.text() == "C:/selected-model/ffum"
    assert page.combined_campaign.text() == "C:/selected-model/combined"
    window.close()
    window.deleteLater()
    qt_app.processEvents()


def test_phase_model_studio_starts_fresh_session_for_unmaterialized_model(
    qt_app: QApplication,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    campaign_roots = {
        name: tmp_path / name for name in ("ffac", "ffum", "combined")
    }
    for root in campaign_roots.values():
        root.mkdir()
        (root / "postflight_manifest.json").write_text("{}", encoding="utf-8")
    session = {
        "campaign_id": "prepared-corrected-frames",
        "ffac_root": str(campaign_roots["ffac"]),
        "ffum_root": str(campaign_roots["ffum"]),
        "combined_root": str(campaign_roots["combined"]),
        "modified_timestamp": 1.0,
    }
    model = {
        "model_id": "phase-model-latest",
        "root": str(tmp_path / "phase-model-latest"),
        "state_label": "Human validation required",
        "created_timestamp": 2.0,
        "heads": [],
    }
    monkeypatch.setattr(
        "darsia.fluidflower.ui.model_training.discover_phase_review_sessions",
        lambda: [session],
    )
    monkeypatch.setattr(
        "darsia.fluidflower.ui.model_training.discover_phase_model_history",
        lambda _root: [],
    )

    window = PhaseModelStudioWindow()
    page = window.training_page
    page.model_history = [dict(model)]
    page.review_sessions = [dict(session)]
    page._populate_validation_models()
    page.reviewer.setText("Reviewer")
    page.review_save_root.setText(str(tmp_path / "reviews"))

    assert page.validation_model.currentData() == model["root"]
    assert page.combined_campaign.text() == str(campaign_roots["combined"].resolve())
    assert page.open_blind_labeling.isEnabled()
    assert page.open_model_correction.isEnabled()
    assert "will generate" in page.validation_model_summary.text()

    launched: list[str] = []
    monkeypatch.setattr(page, "_materialize_and_launch_review", launched.append)
    page.open_blind_labeling.click()
    page.open_model_correction.click()
    assert launched == ["blind_region_labeling", "open_correction"]

    window.close()
    window.deleteLater()
    qt_app.processEvents()


def test_dropdown_values_ignore_wheel_events_application_wide(
    qt_app: QApplication,
) -> None:
    guard = install_interaction_safety(qt_app)

    assert install_interaction_safety(qt_app) is guard
    assert isinstance(guard, ComboBoxWheelGuard)
    combo = QComboBox()
    combo.addItems(["First", "Second"])

    assert guard.eventFilter(combo, QEvent(QEvent.Type.Wheel))
    assert not guard.eventFilter(combo, QEvent(QEvent.Type.FocusIn))
    assert combo.currentIndex() == 0
    combo.show()
    combo.showPopup()
    qt_app.processEvents()
    assert not guard.eventFilter(combo.view().viewport(), QEvent(QEvent.Type.Wheel))
    combo.hidePopup()
    combo.close()


def test_spin_values_require_a_direct_click_before_wheel_input(
    qt_app: QApplication,
) -> None:
    guard = install_interaction_safety(qt_app)
    spin = QSpinBox()
    other = QLabel("Page content")

    assert guard.eventFilter(spin, QEvent(QEvent.Type.Wheel))
    assert not guard.eventFilter(spin, QEvent(QEvent.Type.MouseButtonPress))
    assert not guard.eventFilter(spin, QEvent(QEvent.Type.Wheel))
    assert not guard.eventFilter(other, QEvent(QEvent.Type.MouseButtonPress))
    assert guard.eventFilter(spin, QEvent(QEvent.Type.Wheel))

    spin.close()
    other.close()


def test_home_page_offers_incremental_import_for_selected_project(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Incremental series",
        series_id="incremental-series",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    registry = ProjectRegistry(tmp_path / "registry.json")
    registry.remember(project)
    page = HomePage(registry)
    requested: list[str] = []
    updates: list[str] = []
    page.add_runs_requested.connect(requested.append)
    page.update_run_sources_requested.connect(updates.append)

    assert page.table.columnCount() == 6
    assert not page.add_runs_button.isEnabled()
    page.table.selectRow(0)
    qt_app.processEvents()
    assert page.add_runs_button.isEnabled()
    assert page.update_run_sources_button.text() == "Open selected series setup"
    assert "folder roles" in page.update_run_sources_button.toolTip()
    page.add_runs_button.click()
    page.update_run_sources_button.click()

    assert requested == [str(project.root)]
    assert updates == [str(project.root)]
    page.close()


def test_incremental_source_dialogs_are_drop_in_components(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Dialog series",
        series_id="dialog-series",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    add_dialog = AddRunsDialog(project)
    update_dialog = UpdateRunSourcesDialog(project)

    assert add_dialog.windowTitle() == "Add experiments to Dialog series"
    assert not add_dialog.add_button.isEnabled()
    assert update_dialog.windowTitle() == "Review project setup for Dialog series"
    assert not update_dialog.update_button.isEnabled()
    add_dialog.close()
    update_dialog.close()


def test_add_runs_uses_the_series_image_file_policy(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="RAW series",
        series_id="raw-series",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    project.update_metadata(
        "tools",
        {"source_mapping": {"image_file_policy": "prefer_raw"}},
        updated_by="Reviewer",
    )

    dialog = AddRunsDialog(project)

    assert dialog.sources.image_file_policy_value() == "prefer_raw"
    assert not dialog.sources.image_file_policy.isEnabled()
    assert "global series setting" in dialog.sources.image_file_policy.toolTip()
    dialog.close()


def test_review_setup_exposes_all_runs_rig_class_and_blocks_parent_as_one_run(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Editable series",
        series_id="editable-series",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    mappings: list[dict[str, object]] = []
    for run_id in ("ac17", "ac18"):
        run_root = tmp_path / "raw" / run_id.upper()
        injection = run_root / "CO2 injection"
        injection.mkdir(parents=True)
        frame = injection / "frame001.jpg"
        frame.write_bytes(b"frame")
        mappings.append(
            {
                "run_id": run_id,
                "experiment_folder": str(run_root),
                "experiment_folders": [str(injection)],
                "image_count": 1,
                "baseline_candidate": str(frame),
                "candidate_status": "user_confirmed",
            }
        )
    project.set_runs(
        [{"run_id": row["run_id"], "source_mapping": row} for row in mappings],
        updated_by="Reviewer",
    )

    dialog = UpdateRunSourcesDialog(project)

    assert dialog.editor.sources.scroll_area is None
    assert dialog.editor.content_scroll.widget() is dialog.editor.content_widget
    assert dialog.editor.included_table.minimumHeight() >= 320
    assert dialog.editor.included_table.maximumHeight() > 190
    assert dialog.global_scope_button.isChecked()
    assert not dialog.global_group.isHidden()
    assert dialog.individual_group.isHidden()
    assert dialog.update_button.isHidden()
    assert dialog.run_selector.count() == 2
    assert dialog.run_count_label.text().startswith("2 experiment")
    assert dialog.add_runs_button.text() == "Add experiments..."
    assert dialog.included_table.rowCount() == 2
    assert dialog.included_table.columnCount() == 6
    assert dialog.included_table.item(0, 0).text() == "AC17"
    assert dialog.included_table.item(1, 0).text() == "AC18"
    assert dialog.review_global_structure.text() == (
        "Define and save expected structure..."
    )
    assert dialog.sources.discover_sources.text() == "Re-scan this experiment"
    assert not dialog.sources.structure.isEnabled()
    assert not dialog.apply_rig_button.isEnabled()
    assert not dialog.update_button.isEnabled()
    assert dialog.update_button.text() == "Review changes for AC17..."
    assert "No unsaved individual changes for AC17" in (
        dialog.pending_change_status.text()
    )
    dialog.individual_scope_button.click()
    assert dialog.individual_scope_button.isChecked()
    assert dialog.global_group.isHidden()
    assert not dialog.individual_group.isHidden()
    assert not dialog.update_button.isHidden()
    workspace = ProjectWorkspace(project)
    workspace._review_setup_for_run("ac18")
    qt_app.processEvents()
    assert workspace.navigation.currentItem().data(Qt.ItemDataRole.UserRole) == "setup"
    assert workspace.pages["setup"].run_selector.currentData() == "ac18"
    assert workspace.pages["setup"].individual_scope_button.isChecked()
    workspace.close()
    project.apply_run_layout_template(
        [mappings[0]],
        {
            "example_run_name": "AC17",
            "example_run_root": str(tmp_path / "raw" / "AC17"),
            "folders": [
                {
                    "relative_path": "CO2 injection",
                    "role": "co2_injection_images",
                }
            ],
        },
        updated_by="Reviewer",
    )
    dialog._reload_rows("ac17")
    assert dialog.structure_summary.text().startswith("SAVED r1")
    assert not dialog.update_button.isEnabled()
    dialog.sources.mapping_table.item(0, 4).setText(str(tmp_path / "new-baseline.jpg"))
    qt_app.processEvents()
    assert dialog.update_button.isEnabled()
    dialog.rig_type.setCurrentIndex(dialog.rig_type.findData("large"))
    assert dialog.apply_rig_button.isEnabled()

    dialog._included_run_clicked(1, 0)
    assert dialog.run_selector.currentData() == "ac18"
    assert dialog.sources.source_mapping()[0]["run_id"] == "ac18"
    assert not dialog.update_button.isEnabled()
    assert dialog.update_button.text() == "Review changes for AC18..."

    bad_mapping = dict(mappings[0])
    bad_mapping["experiment_folder"] = str(tmp_path / "raw")
    bad_mapping["experiment_folders"] = [
        str(tmp_path / "raw" / "AC18" / "CO2 injection")
    ]
    issues = _cross_run_mapping_issues(project, "ac17", bad_mapping)
    assert any("contains registered experiment AC18" in issue for issue in issues)
    dialog.close()


def test_global_folder_structure_scope_can_select_all_or_one_run(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    rows: list[dict[str, object]] = []
    for run_id in ("ac17", "ac18", "ac19"):
        run_root = tmp_path / run_id
        injection = run_root / "injection"
        injection.mkdir(parents=True)
        rows.append(
            {
                "run_id": run_id,
                "label": run_id.upper(),
                "source_revision": 1,
                "source_mapping": {
                    "experiment_folder": str(run_root),
                    "experiment_folders": [str(injection)],
                    "image_count": 1,
                },
            }
        )
    template = {
        "example_run_name": "AC17",
        "example_run_root": str(tmp_path / "ac17"),
        "folders": [
            {
                "relative_path": "injection",
                "role": "co2_injection_images",
            }
        ],
    }

    dialog = GlobalRunStructureDialog(
        rows,
        template=template,
        template_revision=2,
        current_run_id="ac18",
    )

    assert dialog.selected_run_ids() == ["ac17", "ac18", "ac19"]
    assert (dialog.representative.currentData() or {})["run_id"] == "ac17"
    assert "3 of 3 selected" == dialog.selection_summary.text()
    dialog._set_selected("current")
    assert dialog.selected_run_ids() == ["ac18"]
    assert "1 of 3 selected" == dialog.selection_summary.text()
    dialog._set_selected("none")
    assert dialog.selected_run_ids() == []
    assert not dialog.continue_button.isEnabled()
    dialog._set_selected("all")
    assert dialog.continue_button.isEnabled()
    dialog.close()


def test_saved_global_structure_example_precedes_current_individual_run(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    rows: list[dict[str, object]] = []
    for run_id in ("ac17", "ac52"):
        root = tmp_path / run_id.upper()
        (root / "injection").mkdir(parents=True)
        rows.append(
            {
                "run_id": run_id,
                "label": run_id.upper(),
                "source_mapping": {
                    "experiment_folder": str(root),
                    "experiment_folders": [str(root / "injection")],
                    "image_count": 1,
                },
            }
        )
    template = {
        "example_run_name": "AC52",
        "example_run_root": str(tmp_path / "AC52"),
        "folders": [
            {
                "relative_path": "injection",
                "role": "co2_injection_images",
            }
        ],
    }

    dialog = GlobalRunStructureDialog(
        rows,
        template=template,
        template_revision=1,
        current_run_id="ac17",
    )

    assert (dialog.representative.currentData() or {})["run_id"] == "ac52"
    assignments, _orders, _tables = _template_editor_values(
        dialog.layout_template(), tmp_path / "AC52"
    )
    assert assignments == {"injection": "co2_injection_images"}
    dialog.close()


def test_integrity_page_selects_only_new_run_sources_after_incremental_import(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Incremental integrity",
        series_id="incremental-integrity",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )

    def mapping(run_id: str) -> dict[str, object]:
        run = tmp_path / "series" / run_id.upper()
        images = run / "CO2 injection"
        images.mkdir(parents=True)
        frame = images / "frame001.jpg"
        frame.write_bytes(b"image")
        return {
            "run_id": run_id,
            "experiment_folder": str(run),
            "experiment_folders": [str(images)],
            "image_count": 1,
            "baseline_candidate": str(frame),
            "clearwater_candidate": "",
            "candidate_status": "user_confirmed",
        }

    first = mapping("ac17")
    project.set_runs(
        [{"run_id": "ac17", "source_mapping": first}],
        updated_by="Reviewer",
    )
    project.set_module_status("data_integrity", "approved", updated_by="Reviewer")
    project.set_run_module_status(
        ["ac17"], "data_integrity", "approved", updated_by="Reviewer"
    )
    project.append_runs([mapping("ac18")], updated_by="Reviewer")
    page = DataIntegrityPage(project)

    selected_rows = [
        page.sources.topLevelItem(row)
        for row in range(page.sources.topLevelItemCount())
        if page.sources.topLevelItem(row).checkState(0) == Qt.CheckState.Checked
    ]
    assert selected_rows
    assert {item.text(2) for item in selected_rows} == {"ac18"}
    assert page._scope_run_ids(page._selected_source_ids()) == ("ac18",)
    assert project.module_states["data_integrity"]["status"] == "ready"
    page.close()


def test_integrity_page_rejects_report_from_an_older_source_revision(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Stale integrity report",
        series_id="stale-integrity-report",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    run = tmp_path / "series" / "AC17"
    first = run / "first"
    second = run / "second"
    first.mkdir(parents=True)
    second.mkdir(parents=True)
    frame1 = first / "frame001.jpg"
    frame2 = second / "frame002.jpg"
    Image.new("RGB", (16, 12), (20, 40, 80)).save(frame1)
    Image.new("RGB", (16, 12), (30, 50, 90)).save(frame2)

    def mapping(folders: list[Path]) -> dict[str, object]:
        return {
            "run_id": "ac17",
            "experiment_folder": str(run),
            "experiment_folders": [str(folder) for folder in folders],
            "image_count": len(folders),
            "baseline_candidate": str(frame1),
            "candidate_status": "user_confirmed",
        }

    project.set_runs(
        [{"run_id": "ac17", "source_mapping": mapping([first])}],
        updated_by="Reviewer",
    )
    report = IntegrityScanner(project).scan(actor="Reviewer")
    IntegrityReportStore(project).save(report, actor="Reviewer")
    project.set_run_module_status(
        ["ac17"], "data_integrity", "approved", updated_by="Reviewer"
    )
    project.set_module_status("data_integrity", "complete", updated_by="Reviewer")
    project.update_run_sources("ac17", mapping([first, second]), updated_by="Reviewer")

    page = DataIntegrityPage(project)

    assert "outdated sources" in page.reports.currentText()
    assert not page.approve.isEnabled()
    assert "Source setup changed" in page.next_step.text()
    selected_run_ids = {
        page.sources.topLevelItem(row).text(2)
        for row in range(page.sources.topLevelItemCount())
        if page.sources.topLevelItem(row).checkState(0) == Qt.CheckState.Checked
    }
    assert selected_run_ids == {"ac17"}
    page.close()


def test_integrity_scan_failure_restores_run_to_ready(
    qt_app: QApplication,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Failed integrity scan",
        series_id="failed-integrity-scan",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    run = tmp_path / "series" / "AC17"
    images = run / "CO2 injection"
    images.mkdir(parents=True)
    frame = images / "frame001.jpg"
    Image.new("RGB", (16, 12), (20, 40, 80)).save(frame)
    project.set_runs(
        [
            {
                "run_id": "ac17",
                "source_mapping": {
                    "run_id": "ac17",
                    "experiment_folder": str(run),
                    "experiment_folders": [str(images)],
                    "image_count": 1,
                    "baseline_candidate": str(frame),
                    "candidate_status": "user_confirmed",
                },
            }
        ],
        updated_by="Reviewer",
    )
    page = DataIntegrityPage(project)
    project.set_run_module_status(
        ["ac17"], "data_integrity", "in_progress", updated_by="Reviewer"
    )
    page._active_scan_run_ids = ("ac17",)
    shown: list[str] = []
    monkeypatch.setattr(
        QMessageBox,
        "critical",
        lambda _parent, _title, message: shown.append(str(message)),
    )

    page._scan_failed("Network source became unavailable")

    state = project.manifest["runs"][0]["workflow"]["data_integrity"]
    assert state["status"] == "ready"
    assert "Network source became unavailable" in state["reason"]
    assert "Network source became unavailable" in page.progress_path.text()
    assert shown == ["Network source became unavailable"]
    page.close()


def test_application_uses_packaged_fluidflower_icon(qt_app: QApplication) -> None:
    icon_path = (
        Path(__file__).parents[2]
        / "src"
        / "darsia"
        / "fluidflower"
        / "ui"
        / "assets"
        / "fluidflower_run10_4h_icon.png"
    )
    assert icon_path.is_file()
    icon = application_icon()
    assert not icon.isNull()
    assert not icon.pixmap(64, 64).isNull()


def test_image_sources_reviews_per_run_hierarchical_mapping(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    run = tmp_path / "runs" / "AC17"
    injection = run / "CO2 injection"
    clearwater = run / "BTB flush"
    overflow = run / "overflow"
    for folder in (injection, clearwater, overflow):
        folder.mkdir(parents=True)
    (injection / "frame01.jpg").write_bytes(b"frame")
    (clearwater / "frame01.jpg").write_bytes(b"clear")
    (overflow / "scale01.jpg").write_bytes(b"scale")

    dialog = RunLayoutDialog(run)
    assert dialog.table.rowCount() == 3
    assert dialog.baseline_rule.currentData() == "first_injection_image"
    assert dialog.clearwater_selector.currentData() == "first"
    template = dialog.template()
    dialog.close()

    page = ImageSourcesPage()
    assert "read-only" in page.raw_data_explanation.text()
    assert "project" in page.raw_data_explanation.text()
    page.experiment_images.setText(str(tmp_path / "runs"))
    page.set_layout_template(template)
    page.set_mapping(
        [
            {
                "run_id": "ac17",
                "experiment_folder": str(run),
                "experiment_folders": [str(injection)],
                "image_count": 1,
                "baseline_candidate": str(injection / "frame01.jpg"),
                "baseline_rule": "first_experiment_image",
                "clearwater_candidate": str(clearwater / "frame01.jpg"),
                "clearwater_folders": [str(clearwater)],
                "clearwater_selector": "first",
                "overflow_image_folders": [str(overflow)],
                "overflow_image_count": 1,
                "overflow_table": "",
                "mapping_status": "ready_for_review",
                "discovery_warnings": [],
                "candidate_status": "requires_user_confirmation",
            }
        ]
    )
    assert not page.isComplete()
    page.mapping_table.item(0, 0).setCheckState(Qt.CheckState.Checked)
    assert page.isComplete()
    mapping = page.source_mapping()[0]
    assert mapping["baseline_candidate"].endswith("frame01.jpg")
    assert mapping["overflow_image_folders"] == [str(overflow)]
    assert mapping["candidate_status"] == "user_confirmed"
    page.close()


def test_image_sources_treats_colorchecker_as_optional_and_buffers_as_chemistry(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    buffer_a = tmp_path / "pH_6.4"
    buffer_b = tmp_path / "pH_8.2"
    alternate = tmp_path / "red dye control"
    for folder in (buffer_a, buffer_b, alternate):
        folder.mkdir()
        (folder / "frame.jpg").write_bytes(b"buffer")

    page = ImageSourcesPage()
    assert page.colorchecker_source.currentData() == "embedded_in_source_images"
    assert not page.colorchecker_images.isEnabled()
    colorchecker_index = page.colorchecker_source.findData("dedicated_folder")
    page.colorchecker_source.setCurrentIndex(colorchecker_index)
    assert page.colorchecker_images.isEnabled()
    assert page.reference_settings()["colorchecker_location"] == "dedicated_folder"

    page.indicator_reference_entries = [
        {
            "kind": "ph_buffer",
            "ph": 6.4,
            "label": "",
            "scope": "series",
            "image_path": str(buffer_a / "frame.jpg"),
        },
        {
            "kind": "ph_buffer",
            "ph": 8.2,
            "label": "",
            "scope": "series",
            "image_path": str(buffer_b / "frame.jpg"),
        },
    ]
    page.alternate_color_reference_images.setText(str(alternate))
    contract = page.indicator_reference_contract()
    assert contract["status"] == "ready"
    assert contract["radiometric_normalization_allowed"] is False
    assert contract["alternate_state_count"] == 1
    assert page.reference_settings()["alternate_color_reference_root"] == str(alternate)
    page.close()


def test_image_correction_owns_model_choice_and_can_outline_detected_chart(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Correction GUI test",
        series_id="correction-gui-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    project.update_metadata(
        "tools",
        {
            "image_reference": {
                "correction_model_selection": "image_correction_module",
                "colorchecker_location": "embedded_in_source_images",
            }
        },
    )
    run_root = tmp_path / "AC52"
    injection = run_root / "CO2 injection"
    injection.mkdir(parents=True)
    baseline = injection / "frame001.jpg"
    last = injection / "frame030.jpg"
    known_stress = injection / "DSC06103.JPG"
    baseline.write_bytes(b"baseline")
    last.write_bytes(b"last")
    known_stress.write_bytes(b"stress")
    project.set_runs(
        [
            {
                "run_id": "ac52",
                "label": "AC52",
                "source_mapping": {
                    "run_id": "ac52",
                    "experiment_folder": str(run_root),
                    "experiment_folders": [str(injection)],
                    "baseline_candidate": str(baseline),
                    "first_image": str(baseline),
                    "last_image": str(last),
                    "co2_sequences": [
                        {
                            "order": 1,
                            "kind": "injection",
                            "folder": str(injection),
                            "first_image": str(baseline),
                            "last_image": str(last),
                        }
                    ],
                    "candidate_status": "user_confirmed",
                },
            }
        ],
        updated_by="Reviewer",
    )
    page = ImageCorrectionPage(project)

    assert page.representative_run.currentData() == "ac52"
    assert page.reference.text() == str(baseline)
    assert page.target.text() == str(last)
    assert page.reference_source_preview.text() == "Preview"
    assert page.target_source_preview.text() == "Preview"
    assert page.reference_source_preview.isEnabled()
    assert page.target_source_preview.isEnabled()
    assert page.images_splitter.widget(0) is page.original_panel
    assert page.left_preview_variant.currentData() == (
        "identity-no-external-color-target-v1"
    )
    assert page.right_preview_variant.currentData() == (
        "uniform24-neutral-tone-affine-v1"
    )
    page.swap_preview_sides.click()
    assert page.left_preview_variant.currentData() == (
        "uniform24-neutral-tone-affine-v1"
    )
    assert page.right_preview_variant.currentData() == (
        "identity-no-external-color-target-v1"
    )
    page.swap_preview_sides.click()
    assert page.images_splitter.widget(0) is page.original_panel
    assert not page.manual_preview_sources.isChecked()
    assert not page.reference.isEnabled()
    assert not page.target.isEnabled()
    assert "apply to every included experiment" in page.scope_explanation.text()
    assert page.batch_use_project_mapping.isChecked()
    assert page.batch_source.isHidden()
    assert "before optical phase interpretation" in page.reference_mode.toolTip()
    assert "embedded in source images" in page.colorchecker_source_status.text()
    assert "before color paths" in page.integration.text()
    assert page.reference_mode.findData("uniform24-neutral-tone-affine-v1") >= 0
    assert page.reference_mode.findData("darsia-baseline-color-correction-v1") >= 0
    labels = [
        page.representative_frame.itemText(index)
        for index in range(page.representative_frame.count())
    ]
    assert any("AC52 9.42 h lighting test" in label for label in labels)
    assert page.preview_frame_slider.maximum() >= 2
    page.original_view.set_array(np.zeros((100, 200, 3), dtype=np.float32))
    page.corrected_view.set_array(np.zeros((100, 200, 3), dtype=np.float32))
    linked_scale = page.corrected_view.transform().m11()
    page.original_view.zoom_requested.emit(1.2)
    assert page.corrected_view.transform().m11() == pytest.approx(
        linked_scale * 1.2
    )
    assert page.show_colorchecker_roi.isChecked()
    assert page.frame_quality.scan_button.text() == "Screen inventoried frames"
    assert page.frame_quality.table.columnCount() == 8
    assert "never exclude" in page.frame_quality.timeline.toolTip().casefold() or (
        "click a marker" in page.frame_quality.timeline.toolTip().casefold()
    )

    view = ZoomableImageView()
    view.set_array(np.zeros((100, 200, 3), dtype=np.float32))
    view.set_roi(
        [[10, 20], [10, 80], [40, 80], [40, 20]],
        source_shape=[100, 200, 3],
    )
    assert view.roi_item.isVisible()
    view.clear_roi()
    assert not view.roi_item.isVisible()
    view.close()
    page.close()


def test_image_source_review_preserves_multiple_overflow_tables(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    run = tmp_path / "AC17"
    injection = run / "CO2 injection"
    injection.mkdir(parents=True)
    image = injection / "frame01.jpg"
    image.write_bytes(b"image")
    table_a = run / "overflow-a.csv"
    table_b = run / "overflow-b.xlsx"
    table_a.write_text("time_h,weight_g\n", encoding="utf-8")
    table_b.write_bytes(b"workbook")
    page = ImageSourcesPage()
    page.set_mapping(
        [
            {
                "run_id": "ac17",
                "experiment_folder": str(run),
                "experiment_folders": [str(injection)],
                "image_count": 1,
                "baseline_candidate": str(image),
                "clearwater_candidate": "",
                "overflow_image_folders": [],
                "overflow_tables": [str(table_a), str(table_b)],
                "overflow_table": str(table_a),
                "mapping_status": "ready_for_review",
                "candidate_status": "user_confirmed",
            }
        ]
    )

    mapping = page.source_mapping()[0]
    assert mapping["overflow_tables"] == [str(table_a), str(table_b)]
    assert mapping["overflow_table"] == str(table_a)
    page.close()


def test_series_run_scan_allows_renaming_and_excluding_folders(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    root = tmp_path / "series"
    for name in ("AC17", "AC18", "documentation"):
        (root / name).mkdir(parents=True)

    dialog = SeriesRunsDialog(root)
    assert dialog.table.rowCount() == 3
    dialog.table.item(0, 1).setText("pilot-17")
    dialog.table.item(2, 0).setCheckState(Qt.CheckState.Unchecked)
    definitions = dialog.definitions()

    assert definitions[0]["run_id"] == "pilot-17"
    assert sum(row["included"] for row in definitions) == 2
    dialog.close()


def test_incremental_run_scan_locks_already_registered_folders(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    root = tmp_path / "series"
    existing = root / "AC17"
    new = root / "AC18"
    existing.mkdir(parents=True)
    new.mkdir()

    dialog = SeriesRunsDialog(
        root,
        reserved_run_ids={"ac17"},
        reserved_folders={str(existing)},
    )

    assert not (dialog.table.item(0, 0).flags() & Qt.ItemFlag.ItemIsUserCheckable)
    assert dialog.table.item(0, 0).checkState() == Qt.CheckState.Unchecked
    assert dialog.table.item(0, 3).text() == "Already in project"
    assert dialog.table.item(1, 0).checkState() == Qt.CheckState.Checked
    assert dialog.definitions()[1]["run_id"] == "ac18"
    dialog.close()


def test_run_layout_bulk_assigns_multiple_folders_to_one_ordered_category(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    run = tmp_path / "Run10"
    for name in ("part-a", "part-b"):
        folder = run / name
        folder.mkdir(parents=True)
        (folder / "frame01.jpg").write_bytes(b"image")

    dialog = RunLayoutDialog(run)
    dialog.table.setRangeSelected(
        QTableWidgetSelectionRange(0, 0, 1, dialog.table.columnCount() - 1),
        True,
    )
    dialog.bulk_role.setCurrentIndex(dialog.bulk_role.findData(ROLE_INJECTION))
    dialog._apply_bulk_role()
    template = dialog.template()

    assert all(role.currentData() == ROLE_INJECTION for role in dialog.role_boxes)
    assert [row["sequence_order"] for row in template["folders"]] == [1, 2]
    dialog.close()


def test_run_layout_shows_and_changes_selected_table_file(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    run = tmp_path / "AC52"
    protocol = run / "protocol"
    protocol.mkdir(parents=True)
    first = protocol / "ac52_notes.xlsx"
    selected = protocol / "ac52_protocol_darsia.xlsx"
    first.write_bytes(b"notes")
    selected.write_bytes(b"protocol")

    dialog = RunLayoutDialog(run)
    row = next(
        index
        for index, record in enumerate(dialog.table_records)
        if record["path"] == str(selected)
    )
    dialog.table_role_boxes[row].setCurrentIndex(
        dialog.table_role_boxes[row].findData(ROLE_PROTOCOL_TABLE)
    )

    assert dialog.table_files.item(row, 1).text().endswith(selected.name)
    assert dialog.table_files.item(row, 1).toolTip() == str(selected)
    table_specs = dialog.template()["table_files"]
    assert (
        next(
            item["role"]
            for item in table_specs
            if item["relative_path"].endswith(selected.name)
        )
        == ROLE_PROTOCOL_TABLE
    )
    dialog.close()


def test_setup_draft_keeps_discovered_overflow_for_later_module(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    store = SetupDraftStore(tmp_path / "drafts")
    run = tmp_path / "series" / "AC17"
    injection = run / "CO2 injection"
    overflow = run / "overflow"
    injection.mkdir(parents=True)
    overflow.mkdir()
    (injection / "frame01.jpg").write_bytes(b"image")
    (overflow / "scale01.jpg").write_bytes(b"scale")

    wizard = NewSeriesWizard(draft_store=store)
    wizard.images.set_run_definitions(
        [{"run_id": "ac17", "folder": str(run), "included": True}]
    )
    wizard.images.set_mapping(
        [
            {
                "run_id": "ac17",
                "experiment_folder": str(run),
                "experiment_folders": [str(injection)],
                "image_count": 1,
                "baseline_candidate": str(injection / "frame01.jpg"),
                "clearwater_candidate": "",
                "overflow_image_folders": [str(overflow)],
                "overflow_image_count": 1,
                "overflow_table": "",
                "mapping_status": "ready_for_review",
                "candidate_status": "user_confirmed",
            }
        ]
    )
    wizard._save_draft()

    values = store.load(wizard.draft_id)["values"]
    assert values["run_definitions"][0]["run_id"] == "ac17"
    assert values["source_mapping"][0]["overflow_image_folders"] == [str(overflow)]
    assert not hasattr(wizard, "overflow")
    wizard.summary.initializePage()
    assert "configure per experiment after Data and integrity" in (
        wizard.summary.summary.toPlainText()
    )
    wizard.close()


def test_incomplete_setup_can_resume_from_saved_page_or_be_discarded(
    qt_app: QApplication,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    store = SetupDraftStore(tmp_path / "drafts")
    store.save(
        {
            "template_id": "ff_ac",
            "name": "Interrupted setup",
            "series_id": "interrupted-setup",
            "parent_folder": str(tmp_path / "projects"),
            "current_page": 2,
        },
        draft_id="draft-1",
    )
    page = HomePage(
        ProjectRegistry(tmp_path / "registry.json"),
        draft_store=store,
    )
    resumed: list[str] = []
    page.resume_setup_requested.connect(resumed.append)

    assert page.draft_table.rowCount() == 1
    assert page.draft_table.item(0, 2).text() == "Apparatus, protocol and sources"
    page.draft_table.selectRow(0)
    page._resume_latest_setup()
    assert resumed == ["draft-1"]

    wizard = NewSeriesWizard(draft_id="draft-1", draft_store=store)
    assert wizard.startId() == 1
    assert wizard.identity.name.text() == "Interrupted setup"
    wizard.show()
    qt_app.processEvents()
    assert wizard.currentId() == 1
    wizard.reject()

    monkeypatch.setattr(
        QMessageBox,
        "question",
        lambda *_args, **_kwargs: QMessageBox.StandardButton.Yes,
    )
    page.refresh()
    page.draft_table.selectRow(0)
    page._discard_selected_setup()
    assert page.draft_table.rowCount() == 0
    assert page.draft_table.isHidden()
    assert (store.root / "discarded" / "draft-1.json").is_file()
    page.close()


def test_open_project_uses_qt_dialog_instead_of_native_shell(
    qt_app: QApplication,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    captured: dict[str, object] = {}

    def choose_directory(*_args: object, **kwargs: object) -> str:
        captured.update(kwargs)
        return ""

    monkeypatch.setattr(QFileDialog, "getExistingDirectory", choose_directory)
    window = FluidFlowerMainWindow(registry=ProjectRegistry(tmp_path / "registry.json"))

    window._open_project_dialog()

    assert captured["options"] == DIRECTORY_DIALOG_OPTIONS
    assert DIRECTORY_DIALOG_OPTIONS & QFileDialog.Option.DontUseNativeDialog
    window.close()


def test_project_package_root_can_be_imported_without_knowing_internal_folder(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    package = tmp_path / "shared-test-package"
    project = FluidFlowerProject.create(
        package / "DarSIA project",
        name="Packaged project",
        series_id="packaged-project",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    window = FluidFlowerMainWindow(registry=ProjectRegistry(tmp_path / "registry.json"))

    window.open_project(package)

    assert window.workspace is not None
    assert window.workspace.project.root == project.root
    assert window.registry.entries()[0]["path"] == str(project.root)
    window.close()


def test_returning_from_home_reuses_project_workspace(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Persistent workspace",
        series_id="persistent-workspace",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    window = FluidFlowerMainWindow(registry=ProjectRegistry(tmp_path / "registry.json"))
    window.show_project(project)
    original_workspace = window.workspace
    original_integrity_page = original_workspace.pages["data_integrity"]

    window.show_home()
    window.open_project(project.root)

    assert window.workspace is original_workspace
    assert window.workspace.pages["data_integrity"] is original_integrity_page
    window.close()


def test_integrity_scan_survives_home_navigation(
    qt_app: QApplication,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    images = tmp_path / "images"
    images.mkdir()
    for index in range(24):
        Image.new("RGB", (32, 24), (index, 40, 80)).save(
            images / f"frame{index:03d}.jpg"
        )
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Background integrity scan",
        series_id="background-integrity-scan",
        template_id="ff_ac",
        rig_type="medium",
        image_sources={"experiment_images": images},
        created_by="Reviewer",
    )
    original_inspect = IntegrityScanner._inspect_entry

    def delayed_inspect(*args: object, **kwargs: object):
        time.sleep(0.02)
        return original_inspect(*args, **kwargs)

    monkeypatch.setattr(IntegrityScanner, "_inspect_entry", delayed_inspect)
    window = FluidFlowerMainWindow(registry=ProjectRegistry(tmp_path / "registry.json"))
    window.show_project(project)
    assert window.workspace is not None
    workspace = window.workspace
    page = workspace.pages["data_integrity"]
    assert isinstance(page, DataIntegrityPage)
    page._start_scan()
    qt_app.processEvents()
    assert page.scan_is_running()

    window.show_home()
    window.open_project(project.root)
    assert window.workspace is workspace
    assert window.workspace.pages["data_integrity"] is page

    deadline = time.monotonic() + 5.0
    while page._thread is not None and time.monotonic() < deadline:
        qt_app.processEvents()
        time.sleep(0.01)
    assert page._thread is None
    assert page.current_report is not None
    assert page.current_report["summary"]["image_count"] == 24
    window.close()


def test_workspace_uses_dedicated_operational_pages(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Workspace test",
        series_id="workspace-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    workspace = ProjectWorkspace(project)
    assert workspace.add_runs_button.text() == "Add experiments"
    assert not hasattr(workspace, "review_setup_button")
    assert isinstance(workspace.pages["setup"], EstablishedSourceSetupEditor)
    assert workspace.page_hosts["setup"] is workspace.pages["setup"]
    assert workspace.pages["setup"].sources.scroll_area is None
    assert isinstance(workspace.pages["data_integrity"], DataIntegrityPage)
    assert isinstance(workspace.pages["geometry_facies"], GeometryFaciesPage)
    assert isinstance(workspace.pages["image_correction"], ImageCorrectionPage)
    assert isinstance(workspace.pages["temporal_sync"], TimelinePage)
    assert isinstance(workspace.pages["optical_calibration"], OpticalCalibrationPage)
    assert isinstance(workspace.pages["phase_interpretation"], PhaseInterpretationPage)
    assert isinstance(workspace.pages["calibration"], CalibrationPage)
    assert isinstance(workspace.pages["mass_analysis"], MassAnalysisPage)
    assert isinstance(workspace.pages["wasserstein"], WassersteinPage)
    assert isinstance(workspace.pages["post_analysis"], PostAnalysisPage)
    assert "reporting" not in workspace.pages
    assert isinstance(workspace.pages["execution"], ExecutionPage)
    assert isinstance(workspace.page_hosts["post_analysis"], QScrollArea)
    assert (
        workspace.page_hosts["post_analysis"].widget()
        is workspace.pages["post_analysis"]
    )
    workspace.close()


def test_workspace_accepts_a_drop_in_page_replacement(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Drop-in page test",
        series_id="drop-in-page-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    registry = default_page_registry(discover_external=False)
    replacement = QLabel("Replacement geometry page")
    registry.register(
        FluidFlowerPagePlugin(
            "geometry_facies",
            lambda _project: replacement,
            scrollable=False,
            source="test",
        ),
        replace_existing=True,
    )

    workspace = ProjectWorkspace(project, page_registry=registry)

    assert workspace.pages["geometry_facies"] is replacement
    assert workspace.page_hosts["geometry_facies"] is replacement
    workspace.close()


def test_geometry_page_accepts_a_drop_in_subpage(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Drop-in subpage test",
        series_id="drop-in-subpage-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    registry = default_subpage_registry(discover_external=False)
    diagnostics = QLabel("External diagnostics")
    registry.register(
        FluidFlowerSubpagePlugin(
            "geometry_facies",
            "external_diagnostics",
            "External diagnostics",
            lambda _parent: diagnostics,
            order=65,
            source="test",
        )
    )

    page = GeometryFaciesPage(project, subpage_registry=registry)

    titles = [page.tabs.tabText(index) for index in range(page.tabs.count())]
    assert titles[-2:] == ["External diagnostics", "QC and history"]
    assert page.tabs.widget(titles.index("External diagnostics")) is diagnostics
    page.close()


def test_geometry_material_examples_autosave_named_boxes(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    image_path = tmp_path / "baseline.png"
    Image.fromarray(np.full((40, 60, 3), (24, 78, 116), dtype=np.uint8)).save(
        image_path
    )
    original = image_path.read_bytes()
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Material examples GUI test",
        series_id="material-examples-gui-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    page = GeometryFaciesPage(project)

    titles = [page.tabs.tabText(index) for index in range(page.tabs.count())]
    assert "Material examples" in titles
    assert page.facies_example_minimum.value() == 8
    page.facies_example_source.setText(str(image_path))
    page.facies_example_minimum.setValue(1)
    page.facies_material_count.setValue(2)
    page._set_facies_material_count()
    assert page.facies_materials.rowCount() == 2

    page.facies_example_material.setCurrentIndex(0)
    page._facies_example_box_drawn(
        {"x": 0.05, "y": 0.05, "width": 0.15, "height": 0.15}
    )
    page.facies_example_material.setCurrentIndex(1)
    page._facies_example_box_drawn(
        {"x": 0.70, "y": 0.65, "width": 0.15, "height": 0.15}
    )

    current = page.facies_example_store.current()
    assert len(current["materials"]) == 2
    assert len(current["examples"]) == 2
    assert page.generate_guided_facies.isEnabled()
    assert page.facies_example_store.history_path.is_file()
    assert image_path.read_bytes() == original
    page.close()


def test_geometry_page_uses_protocol_ids_and_gui_placement_controls(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Injection point GUI test",
        series_id="injection-point-gui-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    page = GeometryFaciesPage(project)

    assert page.tabs.tabText(2) == "Injection points"
    page._add_injector()
    assert page.injector_table.item(0, 0).text() == "ip1"
    assert page.injector_table.item(0, 1).text() == "Injection Point 1"
    assert page.preview_injector_reference.text() == "Preview placement image"
    assert page.place_injector.text() == "Place selected injection point"

    page._populate_injectors(
        [
            {
                "injector_id": "ip1",
                "protocol_location_id": "ip1",
                "name": "Injection Point 1",
                "x_m": 0.30,
                "y_m": 0.10,
                "protocol_x_m": 0.30,
                "protocol_y_m": 0.10,
                "coordinate_source": "protocol:protocol.xlsx",
            }
        ]
    )
    page.injector_table.item(0, 2).setText("0.35")
    assert page.injector_table.item(0, 5).text() == "manual_coordinate_entry"
    assert "GUI and Excel positions differ" in page.injector_position_status.text()
    collected = page._collect_config()["injectors"][0]
    assert collected["x_m"] == 0.35
    assert collected["protocol_x_m"] == 0.30
    page.close()


def test_geometry_page_loads_excel_points_and_preserves_manual_position(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    protocol = write_protocol_workbook_template(
        tmp_path / "protocol.xlsx",
        layout="single_run",
        run_ids=("ac52",),
    )
    workbook = load_workbook(protocol)
    sheet = workbook["Injection points"]
    headers = [cell.value for cell in sheet[1]]
    x_column = headers.index("x_m") + 1
    y_column = headers.index("y_m") + 1
    sheet.cell(2, x_column, 0.30)
    sheet.cell(2, y_column, 0.10)
    workbook.save(protocol)

    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Excel injection-point GUI test",
        series_id="excel-injection-point-gui-test",
        template_id="ff_ac",
        rig_type="medium",
        protocol_path=protocol,
        created_by="Reviewer",
    )
    page = GeometryFaciesPage(project)

    assert page.protocol_path.text() == str(protocol.resolve())
    assert page.load_protocol_points.text() == "Load / update from Excel protocol"
    page._load_injection_points_from_protocol()
    assert page.injector_table.rowCount() == 2
    assert page.injector_table.item(0, 0).text() == "ip1"
    assert page.injector_table.item(0, 1).text() == "Injection Point 1"
    assert float(page.injector_table.item(0, 2).text()) == pytest.approx(0.3)

    page.injector_table.item(0, 2).setText("0.35")
    assert page.injector_table.item(0, 5).text() == "manual_coordinate_entry"
    workbook = load_workbook(protocol)
    sheet = workbook["Injection points"]
    sheet.cell(2, x_column, 0.32)
    workbook.save(protocol)
    page._load_injection_points_from_protocol()

    collected = page._collect_config()["injectors"][0]
    assert collected["x_m"] == 0.35
    assert collected["protocol_x_m"] == 0.32
    assert "GUI and Excel positions differ" in page.injector_position_status.text()
    page.close()


def test_workspace_context_help_tracks_module_and_subsection(
    qt_app: QApplication,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Context help test",
        series_id="context-help-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    workspace = ProjectWorkspace(project)

    assert all(definition.help_text for definition in MODULE_DEFINITIONS)
    assert workspace.current_section_label.text() == "Workflow overview"

    workspace.select_module("geometry_facies")
    qt_app.processEvents()
    assert workspace.current_section_label.text() == "Geometry and facies / Inputs"

    geometry_page = workspace.pages["geometry_facies"]
    layers_index = next(
        index
        for index in range(geometry_page.tabs.count())
        if geometry_page.tabs.tabText(index) == "Layers"
    )
    geometry_page.tabs.setCurrentIndex(layers_index)
    qt_app.processEvents()
    assert workspace.current_section_label.text() == "Geometry and facies / Layers"
    assert workspace.context_help_button.toolTip() == (
        "About Geometry and facies / Layers"
    )

    captured: dict[str, str] = {}

    def show_information(_parent: object, title: str, body: str) -> None:
        captured.update(title=title, body=body)

    monkeypatch.setattr(QMessageBox, "information", show_information)
    workspace._show_context_help()
    assert captured["title"] == "Geometry and facies / Layers"
    assert "porosity" in captured["body"]
    workspace.close()


def test_execution_page_exposes_hardware_queue_and_worker_controls(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Execution GUI test",
        series_id="execution-gui-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    page = ExecutionPage(project)

    assert page.execution_queue_mode.currentData() == "local"
    assert page.execution_queue_path.text() == str(
        project.root / "cache" / "execution" / "queue"
    )
    assert page.execution_result_path.text() == str(
        project.root / "results" / "calibration"
    )
    assert page.execution_module.currentData() == "calibration"
    assert page.execution_adapter.currentData() == "plan-verification-v1"
    assert page.execution_backend.currentData() == "cpu-float64"
    assert page.execution_workers.value() == 1
    assert page.execution_adapters.rowCount() >= 1
    page.execution_module.setCurrentIndex(
        page.execution_module.findData("post_analysis")
    )
    assert page.execution_adapter.currentData() == "post-analysis-v1"
    assert page.execution_result_path.text() == str(
        project.root / "results" / "post-analysis"
    )
    custom_post = tmp_path / "large-results" / "post"
    page.execution_result_path.setText(str(custom_post))
    page.execution_module.setCurrentIndex(page.execution_module.findData("wasserstein"))
    assert page.execution_adapter.currentData() == "wasserstein-cpu-float64-v1"
    assert page.execution_result_path.text() == str(
        project.root / "results" / "wasserstein"
    )
    page.execution_module.setCurrentIndex(
        page.execution_module.findData("post_analysis")
    )
    assert page.execution_result_path.text() == str(custom_post.resolve())
    page._persist_execution_paths()
    stored = project.manifest["tools"]["execution"]
    assert stored["default_result_root"] == str(project.root / "results")
    assert stored["module_result_roots"]["post_analysis"] == str(custom_post.resolve())
    page.execution_reset_result_path.click()
    assert page.execution_result_path.text() == str(
        project.root / "results" / "post-analysis"
    )
    assert page.execution_module.findData("reporting") == -1
    assert not hasattr(page, "shell_command")
    page.close()


def test_execution_page_migrates_legacy_result_root_per_module(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Legacy execution storage",
        series_id="legacy-execution-storage",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    legacy_root = tmp_path / "old-execution-results"
    project.update_metadata(
        "tools",
        {"execution": {"result_root": str(legacy_root)}},
        updated_by="Reviewer",
    )

    page = ExecutionPage(project)
    assert page.execution_result_path.text() == str(
        legacy_root.resolve() / "calibration"
    )
    page.execution_module.setCurrentIndex(
        page.execution_module.findData("mass_analysis")
    )
    assert page.execution_result_path.text() == str(
        legacy_root.resolve() / "mass-analysis"
    )
    page.execution_reset_result_path.click()
    assert page.execution_result_path.text() == str(
        project.root / "results" / "mass-analysis"
    )
    page.execution_module.setCurrentIndex(page.execution_module.findData("calibration"))
    assert page.execution_result_path.text() == str(
        legacy_root.resolve() / "calibration"
    )
    page.close()


def test_phase_page_registers_shared_model_and_exposes_revision_controls(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Phase GUI test",
        series_id="phase-gui-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    page = PhaseInterpretationPage(project)

    assert page.phase_model_bundle.text() == ""
    assert page.inspect_phase_model_button.isEnabled()
    assert page.save_phase_model_plan.isEnabled()
    assert page.phase_campaign_manifest.text() == ""
    assert page.phase_review_package.text() == ""
    assert page.phase_candidates.columnCount() == 6
    assert page.phase_selected.columnCount() == 3
    assert page.save_phase_revision.isEnabled()
    assert page.phase_workflow_tabs.count() == 2
    assert page.phase_workflow_tabs.tabText(0) == "Active learning review"
    assert page.phase_workflow_tabs.tabText(1) == "Model validation"
    assert page.open_active_learning_review.isEnabled()
    assert page.open_model_validation.isEnabled()
    assert not hasattr(page, "train_phase_model")
    assert any(
        "complete corrected 2D image sequence" in label.text()
        for label in page.phase_workflow_tabs.widget(1).findChildren(QLabel)
    )
    page.close()


def test_calibration_page_exposes_plan_and_separate_result_review(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Calibration GUI test",
        series_id="calibration-gui-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    page = CalibrationPage(project)

    assert page.calibration_trials.value() == 5000
    assert page.calibration_backend.currentData() == "cpu-float64"
    assert page.calibration_dtype.text() == "float64"
    assert page.calibration_resolution_policy.currentData() == "native_registered_grid"
    assert page.calibration_backend.findData("cpu-float32") >= 0
    assert page.calibration_backend.findData("cuda-float64") >= 0
    assert page.calibration_backend.findData("opencl-float64") >= 0
    assert page.calibration_same_objective.isChecked()
    assert not page.calibration_same_objective.isEnabled()
    assert page.calibration_indicator.text() == "BTB"
    assert page.calibration_parameter_table.rowCount() == 11
    assert "diffusion_coefficient_m2_s" in page.calibration_parameter_widgets
    assert not page.calibration_parameter_widgets["diffusion_coefficient_m2_s"][
        0
    ].isChecked()
    assert page.calibration_evaluation_manifest.text() == ""
    assert page.calibration_result_manifest.text() == ""
    assert "generic optical branch" in page.indicator_response_status.text()
    assert page.calibration_run_review_table.columnCount() == 6
    assert "No mass-calibration result" in page.calibration_result_review_summary.text()
    page.calibration_backend.setCurrentIndex(
        page.calibration_backend.findData("cpu-float32")
    )
    assert page.calibration_dtype.text() == "float32"
    assert "cpu-float32" in page.calibration_quality_basis.text()
    assert "native registered grid" in page.calibration_quality_basis.text()
    assert not hasattr(page, "run_optuna")
    assert not hasattr(page, "train_model")
    page.close()


def test_mass_page_is_task_planning_only(qt_app: QApplication, tmp_path: Path) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Mass GUI test",
        series_id="mass-gui-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    page = MassAnalysisPage(project)

    assert page.mass_maximum_images.value() == 100
    assert page.mass_drift_tolerance.value() == 5.0
    assert page.mass_backend.currentData() == "cpu-float64"
    assert page.mass_dtype.text() == "float64"
    assert page.mass_resolution_policy.currentData() == "native_registered_grid"
    assert page.mass_backend.findData("cpu-float32") >= 0
    assert page.mass_backend.findData("cuda-float64") >= 0
    assert not page.mass_wasserstein_cache.isChecked()
    assert not page.mass_wasserstein_cell_size.isEnabled()
    assert page.mass_products["mass_total"].isChecked()
    assert page.mass_products["integrated_mass_csv"].isChecked()
    assert page.mass_optical_branch.currentData() is None
    page.mass_backend.setCurrentIndex(page.mass_backend.findData("cpu-float32"))
    page.mass_resolution_policy.setCurrentIndex(
        page.mass_resolution_policy.findData("scale_factor")
    )
    page.mass_scale_percent.setValue(50.0)
    config = page._mass_configuration()
    assert page.mass_dtype.text() == "float32"
    assert "50% of registered grid" in page.mass_quality_basis.text()
    assert config["execution_profile"]["primary_backend"] == "cpu-float32"
    assert config["exports"]["numerical_dtype"] == "float32"
    assert config["quality_profile"]["numerical_resolution"]["scale_factor"] == 0.5
    assert not config["optical_branch"]["selection_confirmed"]
    assert not hasattr(page, "run_mass_analysis")
    page.close()


def test_calibration_and_mass_quality_choices_are_stage_local(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Independent quality GUI test",
        series_id="independent-quality-gui-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    calibration = CalibrationPage(project)
    mass = MassAnalysisPage(project)

    calibration.calibration_backend.setCurrentIndex(
        calibration.calibration_backend.findData("cpu-float64")
    )
    calibration.calibration_preview_dimension.setValue(2560)
    mass.mass_backend.setCurrentIndex(mass.mass_backend.findData("cpu-float32"))
    mass.mass_resolution_policy.setCurrentIndex(
        mass.mass_resolution_policy.findData("scale_factor")
    )
    mass.mass_scale_percent.setValue(50.0)
    mass.mass_preview_dimension.setValue(1024)

    calibration_config = calibration._calibration_configuration()
    mass_config = mass._mass_configuration()

    assert calibration_config["execution_profile"]["primary_backend"] == ("cpu-float64")
    assert (
        calibration_config["quality_profile"]["numerical_resolution"]["policy"]
        == "native_registered_grid"
    )
    assert (
        calibration_config["quality_profile"]["preview"]["maximum_dimension_px"] == 2560
    )
    assert mass_config["execution_profile"]["primary_backend"] == "cpu-float32"
    assert mass_config["quality_profile"]["numerical_resolution"]["scale_factor"] == 0.5
    assert mass_config["quality_profile"]["preview"]["maximum_dimension_px"] == 1024

    calibration.close()
    mass.close()


def test_wasserstein_page_is_task_planning_only(
    qt_app: QApplication, tmp_path: Path
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Wasserstein GUI test",
        series_id="wasserstein-gui-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    page = WassersteinPage(project)

    assert page.wasserstein_backend.text() == "cpu-float64"
    assert page.wasserstein_workers.value() == 12
    assert page.wasserstein_cache_timepoints.value() == 2
    assert not page.wasserstein_physical_resolution.isChecked()
    assert not page.wasserstein_cell_size.isEnabled()
    assert all(
        checkbox.isChecked() for checkbox in page.wasserstein_components.values()
    )
    assert all(checkbox.isChecked() for checkbox in page.wasserstein_domains.values())
    assert not hasattr(page, "run_wasserstein")
    page.close()


def test_post_analysis_page_is_task_planning_only(
    qt_app: QApplication, tmp_path: Path
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Post-analysis GUI test",
        series_id="post-analysis-gui-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    page = PostAnalysisPage(project)

    assert page.post_workers.value() == 24
    assert page.post_encoder.currentData() == "auto-hardware-if-available"
    assert page.post_movie_width.value() == 1600
    assert page.post_movie_fps.value() == 6.0
    assert page.post_pre_injection.isChecked()
    assert not page.post_pre_injection.isEnabled()
    assert not page.post_products["phase_map_movie_black_background"].isChecked()
    assert all(
        page.post_products[key].isChecked() and not page.post_products[key].isEnabled()
        for key in (
            "integrated_mass_csv",
            "per_frame_diagnostics_csv",
            "provenance_json",
        )
    )
    assert not hasattr(page, "run_post_analysis")
    page.close()


def test_geometry_page_exposes_portable_facies_dataset_controls(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Facies export GUI test",
        series_id="facies-export-gui",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    page = GeometryFaciesPage(project)

    assert page.tabs.indexOf(page.training_data_tab) >= 0
    assert page.facies_group_id.text() == "facies-export-gui"
    assert page.facies_split.count() == 4
    assert page.facies_split.itemData(2) == "holdout"
    assert page.facies_split.itemData(3) == "pool"
    assert page.facies_dataset_root.text() == str(
        project.root / "annotations" / "facies"
    )
    assert "0 bundles" in page.facies_dataset_summary.text()
    assert page.facies_model_dataset_roots.text() == str(
        project.root / "annotations" / "facies"
    )
    assert page.facies_model_root.text() == str(
        project.root / "results" / "facies_models"
    )
    assert page.facies_model_split_plan.text() == ""
    assert page.facies_model_estimators.value() == 128
    assert page.facies_model_pixels.value() == 120_000
    assert page.facies_model_calibrate.isChecked()
    assert not page.facies_model_position.isChecked()
    assert page.train_facies_model.isEnabled()
    page.close()


def test_integrity_page_loads_current_report_and_allows_clean_review(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    images = tmp_path / "images"
    images.mkdir()
    Image.new("RGB", (32, 24), (20, 40, 80)).save(images / "frame001.jpg")
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Integrity GUI test",
        series_id="integrity-gui-test",
        template_id="ff_ac",
        rig_type="medium",
        image_sources={"experiment_images": images},
        created_by="Reviewer",
    )
    report = IntegrityScanner(project).scan(actor="Reviewer")
    IntegrityReportStore(project).save(report, actor="Reviewer")
    page = DataIntegrityPage(project)
    assert page.current_report is not None
    assert "1 files" in page.report_summary.text()
    assert page.approve.isEnabled()
    page.close()


def test_integrity_page_explains_scan_and_worker_choices(
    qt_app: QApplication,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from darsia.fluidflower.pages.data_integrity import page as integrity_page

    monkeypatch.setattr(
        integrity_page,
        "detect_host_resources",
        lambda: {
            "cpu": {"logical_cores": 32, "physical_cores": 16},
            "memory": {
                "total_bytes": 256 * 1024**3,
                "available_bytes": 256 * 1024**3,
            },
        },
    )
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Integrity help test",
        series_id="integrity-help-test",
        template_id="ff_um",
        rig_type="large",
        created_by="Reviewer",
    )
    page = DataIntegrityPage(project)

    assert "first check" in str(page.mode.itemData(0, Qt.ItemDataRole.ToolTipRole))
    assert "Every image is fully read" in str(
        page.mode.itemData(1, Qt.ItemDataRole.ToolTipRole)
    )
    assert "Use 3 normally" in page.retry_count.toolTip()
    assert page.parallel_reads.value() == 8
    assert page.parallel_reads.maximum() == 64
    assert "32 logical CPU cores" in page.parallel_reads.toolTip()
    assert "editable during a scan" in page.parallel_reads.toolTip()
    assert "not analysis accuracy" in page.parallel_reads.toolTip()
    page.close()


def test_integrity_page_groups_findings_and_identifies_the_run(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Integrity finding presentation",
        series_id="integrity-finding-presentation",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    run = tmp_path / "series" / "AC17"
    images = run / "CO2 injection"
    images.mkdir(parents=True)
    baseline = images / "frame001.jpg"
    Image.new("RGB", (32, 24), (20, 40, 80)).save(baseline)
    project.set_runs(
        [
            {
                "run_id": "ac17",
                "source_mapping": {
                    "run_id": "ac17",
                    "experiment_folder": str(run),
                    "experiment_folders": [str(images)],
                    "image_count": 1,
                    "baseline_candidate": str(baseline),
                    "candidate_status": "user_confirmed",
                },
            }
        ],
        updated_by="Reviewer",
    )
    source = next(
        row
        for row in project.current_source_records()
        if row.get("kind") == "co2_image_folder"
    )
    source_id = str(source["source_id"])
    page = DataIntegrityPage(project)
    page.current_report = {
        "summary": {
            "file_count": 10,
            "image_count": 10,
            "error_count": 0,
            "warning_count": 2,
            "read_retry_count": 0,
        },
        "selected_source_ids": [source_id],
        "sources": [
            {
                "source_id": source_id,
                "kind": "co2_image_folder",
                "run_id": "ac17",
                "path": str(images),
                "available": True,
                "file_count": 10,
                "image_count": 10,
                "total_size_bytes": 1024,
                "issue_count": 2,
            }
        ],
        "issues": [
            {
                "severity": "warning",
                "code": "missing_sequence_numbers",
                "source_id": source_id,
                "path": str(images),
                "message": "One image number is missing.",
            },
            {
                "severity": "warning",
                "code": "missing_sequence_numbers",
                "source_id": source_id,
                "path": str(images),
                "message": "Two image numbers are missing.",
            },
        ],
    }
    page._populate_report()

    assert page.sources.columnCount() == 10
    assert not hasattr(page, "source_summary")
    assert page.findings.rowCount() == 1
    assert page.findings.item(0, 2).text() == "AC17"
    assert page.findings.item(0, 3).text() == "CO2 image folder"
    assert page.findings.item(0, 4).text() == "2"
    assert page.start.isHidden()
    assert page.review.isHidden()
    assert page.complete.isHidden()

    requested_runs: list[str] = []
    page.review_setup_requested.connect(requested_runs.append)
    page.findings.selectRow(0)
    page._request_source_review()
    assert requested_runs == ["ac17"]
    assert "Next action" in page.finding_guidance.text()
    page.close()


def test_approving_all_run_integrity_reports_completes_the_series_step(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Complete run integrity",
        series_id="complete-run-integrity",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    run = tmp_path / "series" / "AC17"
    images = run / "CO2 injection"
    images.mkdir(parents=True)
    baseline = images / "frame001.jpg"
    Image.new("RGB", (32, 24), (20, 40, 80)).save(baseline)
    project.set_runs(
        [
            {
                "run_id": "ac17",
                "source_mapping": {
                    "run_id": "ac17",
                    "experiment_folder": str(run),
                    "experiment_folders": [str(images)],
                    "image_count": 1,
                    "baseline_candidate": str(baseline),
                    "candidate_status": "user_confirmed",
                },
            }
        ],
        updated_by="Reviewer",
    )
    report = IntegrityScanner(project).scan(actor="Reviewer")
    IntegrityReportStore(project).save(report, actor="Reviewer")
    project.set_run_module_status(
        ["ac17"],
        "data_integrity",
        "review_required",
        updated_by="Reviewer",
    )
    page = DataIntegrityPage(project)

    page._set_status("approved")

    assert project.module_states["data_integrity"]["status"] == "complete"
    run_state = project.manifest["runs"][0]["workflow"]["data_integrity"]
    assert run_state["status"] == "approved"
    assert page.continue_button.isVisibleTo(page)
    assert "is complete" in page.next_step.text()
    requested_modules: list[str] = []
    page.module_requested.connect(requested_modules.append)
    page.continue_button.click()
    assert requested_modules == ["overflow"]
    page.close()


def test_integrity_gate_uses_latest_report_when_viewing_history(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    images = tmp_path / "images"
    images.mkdir()
    Image.new("RGB", (32, 24), (20, 40, 80)).save(images / "frame001.jpg")
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Integrity report history",
        series_id="integrity-report-history",
        template_id="ff_ac",
        rig_type="medium",
        image_sources={"experiment_images": images},
        created_by="Reviewer",
    )
    store = IntegrityReportStore(project)
    clean_report = IntegrityScanner(project).scan(actor="Reviewer")
    store.save(clean_report, actor="Reviewer")
    images.rename(tmp_path / "images-unavailable")
    failed_report = IntegrityScanner(project).scan(actor="Reviewer")
    store.save(failed_report, actor="Reviewer")

    page = DataIntegrityPage(project)
    old_index = page.reports.findData(clean_report["scan_id"])
    page.reports.setCurrentIndex(old_index)
    assert page.current_report is not None
    assert page.current_report["scan_id"] == clean_report["scan_id"]
    assert not page.approve.isEnabled()
    page.close()


def test_geometry_page_round_trips_validated_config_and_respects_gate(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    baseline = tmp_path / "baseline.png"
    Image.new("RGB", (8, 6), (20, 40, 80)).save(baseline)
    depth = tmp_path / "depth.npy"
    labels = tmp_path / "labels.npy"
    np.save(depth, np.full((6, 8), 0.010123456789, dtype=np.float64))
    np.save(labels, np.ones((6, 8), dtype=np.int16))
    protocol = tmp_path / "protocol.csv"
    protocol.write_text("location_x,location_y\n0.45,0.03\n", encoding="utf-8")
    config = {
        "schema": 1,
        "format": "darsia-fluidflower-geometry-config-v1",
        "rig_type": "medium",
        "source_rig_path": "",
        "baseline_path": str(baseline),
        "depth_path": str(depth),
        "labels_path": str(labels),
        "facies_source_path": "",
        "facies_properties_source_path": "",
        "protocol_source_path": str(protocol),
        "width_m": 0.9,
        "height_m": 0.49,
        "coordinate_convention": "origin_lower_left_x_right_y_up",
        "grouping_mode": "shared_facies",
        "permeability_unit": "mD",
        "regions": [
            {
                "label_id": 1,
                "facies_id": 1,
                "name": "Sand",
                "analysis_role": "porous_medium",
                "porosity": 0.4123456789012345,
                "permeability": 473.1234567890123,
                "enabled": True,
            }
        ],
        "injectors": [
            {
                "injector_id": "injector-1",
                "name": "Injector 1",
                "x_m": 0.45,
                "y_m": 0.03,
                "uncertainty_m": None,
                "coordinate_source": "protocol",
            }
        ],
    }
    report = validate_geometry_config(config)
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Geometry GUI test",
        series_id="geometry-gui-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    page = GeometryFaciesPage(project)
    page._populate_config(config)
    page._show_validation(report)

    assert page.run_geometry_policy.currentData() == "allow_reviewed_run_overrides"
    assert page.run_geometry_change_detection.isChecked()
    assert page.run_geometry_displacement_threshold.value() == pytest.approx(0.5)
    assert page.run_geometry_affected_threshold.value() == pytest.approx(10.0)
    assert geometry_config_digest(page._collect_config()) == report["config_sha256"]
    assert page.save_button.isEnabled()
    assert not page.approve.isEnabled()

    project.set_module_status("data_integrity", "approved")
    project.set_module_status("image_correction", "approved")
    revision = GeometryRevisionStore(project).save_revision(
        config,
        actor="Reviewer",
        reason="GUI gate test",
        validation=report,
    )
    project.set_module_status("geometry_facies", "review_required")
    page.refresh()
    assert revision["revision_id"]
    assert page.approve.isEnabled()
    page.close()


def test_geometry_page_materializes_autosaved_boundary_edits(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    baseline = tmp_path / "baseline.png"
    Image.new("RGB", (8, 6), (20, 40, 80)).save(baseline)
    depth = tmp_path / "depth.npy"
    labels = tmp_path / "labels.npy"
    np.save(depth, np.full((6, 8), 0.01, dtype=np.float64))
    np.save(labels, np.ones((6, 8), dtype=np.int16))
    config = {
        "schema": 1,
        "format": "darsia-fluidflower-geometry-config-v1",
        "rig_type": "medium",
        "baseline_path": str(baseline),
        "depth_path": str(depth),
        "labels_path": str(labels),
        "width_m": 0.9,
        "height_m": 0.49,
        "grouping_mode": "shared_facies",
        "regions": [
            {
                "label_id": 1,
                "facies_id": 1,
                "name": "Sand",
                "analysis_role": "porous_medium",
                "porosity": 0.4,
                "permeability": 473.0,
                "enabled": True,
            }
        ],
        "injectors": [],
    }
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Geometry boundary GUI test",
        series_id="geometry-boundary-gui-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    page = GeometryFaciesPage(project)
    page._populate_config(config)
    page.preview_data = load_geometry_preview(config)
    page.canvas.set_model(page.preview_data, config)
    page._add_edit_target()
    page._label_edit_drawn("polygon", [[0.0, 0.0], [0.6, 0.0], [0.6, 0.6], [0.0, 0.6]])

    assert page.label_session is not None
    assert page.label_session.events_path.is_file()
    assert page.materialize_label_edits.isEnabled()
    assert not page.validate_button.isEnabled()

    page._materialize_label_edits()
    materialized = Path(page.labels_path.text())
    with np.load(materialized, allow_pickle=False) as archive:
        assert 2 in np.unique(archive["array"])
    assert np.array_equal(np.load(labels), np.ones((6, 8), dtype=np.int16))
    assert page.validate_button.isEnabled()
    assert not page.materialize_label_edits.isEnabled()
    page.close()


def test_timeline_page_displays_raw_and_effective_values(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Timeline GUI test",
        series_id="timeline-gui-test",
        template_id="ff_um",
        rig_type="large",
        created_by="Reviewer",
    )
    page = TimelinePage(project)
    event = page.store.add_event(
        run_id="run18",
        source="images",
        event_type="image_capture",
        time_h=1.0,
        actor="Reviewer",
    )
    page.store.set_override(
        event["event_id"],
        {"time_h": 1.25},
        actor="Reviewer",
        reason="Clock alignment",
    )
    page.selected_event_id = event["event_id"]
    page.refresh()
    group = page.tree.topLevelItem(0)
    row = group.child(0)
    assert row.text(1) == "1.0000"
    assert row.text(3) == "1.2500"
    assert "time_h" in row.text(2)
    page.close()


def test_timeline_page_filters_one_run_and_uses_domain_streams(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Timeline river GUI test",
        series_id="timeline-river-gui-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    store = TimelineStore(project)
    store.add_event(
        run_id="ac52",
        source="protocol",
        event_type="camera_interval_start",
        time_h=0.0,
        actor="Reviewer",
    )
    store.add_event(
        run_id="ac52",
        source="protocol",
        event_type="injection_steady_start",
        time_h=1.0,
        actor="Reviewer",
    )
    store.add_event(
        run_id="ac53",
        source="scale",
        event_type="overflow_measurement",
        time_h=2.0,
        actor="Reviewer",
    )

    page = TimelinePage(project)

    assert page.run_filter.count() == 2
    assert page.run_filter.currentData() == "ac52"
    assert {event["run_id"] for event in page.plot.events} == {"ac52"}
    assert {
        page.plot.stream_key(event) for event in page.plot.events
    } == {"camera", "injection"}

    page.run_filter.setCurrentIndex(page.run_filter.findData("ac53"))
    assert {event["run_id"] for event in page.plot.events} == {"ac53"}
    assert page.plot.stream_key(page.plot.events[0]) == "overflow"
    assert page.timeline_summary.text().startswith("AC53: 1 events")
    page.close()


def test_timeline_selection_is_linked_and_timeline_is_zoomable(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Timeline interaction GUI test",
        series_id="timeline-interaction-gui-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    store = TimelineStore(project)
    event = store.add_event(
        run_id="ac52",
        source="camera",
        event_type="camera_interval_start",
        time_h=0.0,
        value=30.0,
        unit="s",
        path="DSC00121.JPG",
        actor="Reviewer",
    )
    injection = store.add_event(
        run_id="ac52",
        source="injection",
        event_type="injection_steady_start",
        time_h=1.0,
        actor="Reviewer",
    )
    page = TimelinePage(project)
    page.resize(1200, 850)
    page.show()
    qt_app.processEvents()

    page.plot.event_selected.emit(injection["event_id"])
    qt_app.processEvents()

    assert page.selected_event_id == injection["event_id"]
    assert page.plot.selected_event_id == injection["event_id"]
    assert page.tree.currentItem().data(0, Qt.ItemDataRole.UserRole) == injection[
        "event_id"
    ]
    selected_color = page.tree.currentItem().background(0).color()
    assert selected_color == page.tree.currentItem().background(1).color()

    page.timeline_zoom.setValue(100)
    qt_app.processEvents()
    assert page.timeline_scroll.zoom_factor == 10.0
    assert page.plot.width() > page.timeline_scroll.viewport().width()
    assert page.timeline_scroll.horizontalScrollBar().maximum() > 0

    page.timeline_zoom.setValue(300)
    qt_app.processEvents()
    assert page.timeline_scroll.zoom_factor == 1000.0
    assert page.timeline_zoom_label.text() == "1,000x"

    exact, exact_basis = page._linked_image(page.plot.events, event)
    inferred, inferred_basis = page._linked_image(page.plot.events, injection)
    assert exact == "DSC00121.JPG"
    assert "Exact image" in exact_basis
    assert inferred == "~DSC00241.JPG"
    assert "Expected image" in inferred_basis
    page.close()


def test_timeline_page_adds_exact_camera_end_from_integrity_inventory(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Timeline camera endpoint test",
        series_id="timeline-camera-endpoint-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    store = TimelineStore(project)
    store.add_event(
        run_id="ac52",
        source="camera",
        event_type="camera_interval_start",
        time_h=0.0,
        value=30.0,
        unit="s",
        path="DSC00121.JPG",
        actor="Reviewer",
    )
    endpoints = [
        {
            "source_id": "camera-source",
            "run_id": "ac52",
            "source_path": str(tmp_path / "images"),
            "image_count": 121,
            "last": {"path": str(tmp_path / "images" / "DSC00241.JPG")},
        }
    ]

    events = [*store.effective_events(), *TimelinePage._camera_end_events(store.effective_events(), endpoints)]
    endpoint = next(row for row in events if row["event_type"] == "camera_acquisition_end")

    assert endpoint["time_h"] == pytest.approx(1.0)
    assert Path(endpoint["path"]).name == "DSC00241.JPG"
    assert endpoint["label"] == "Camera acquisition ended at DSC00241.JPG"
    assert endpoint["metadata"]["read_only_derived"] is True

    page = TimelinePage(project)
    page.protocol_sync_result["camera_sequence_endpoints"] = endpoints
    page.refresh()
    assert any(
        row["event_type"] == "camera_acquisition_end" for row in page.plot.events
    )
    page.close()


def test_timeline_page_reads_registered_protocol_and_has_explicit_run_approval(
    qt_app: QApplication,
    tmp_path: Path,
) -> None:
    workbook = write_protocol_workbook_template(
        tmp_path / "ac52_protocol.xlsx",
        layout="single_run",
        run_ids=("ac52",),
    )
    project = FluidFlowerProject.create(
        tmp_path / "project",
        name="Registered protocol GUI test",
        series_id="registered-protocol-gui-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )
    project.manifest["runs"] = [
        {
            "run_id": "ac52",
            "source_mapping": {"protocol_table": str(workbook)},
            "workflow": project.module_states,
        }
    ]
    project._write_manifest()

    page = TimelinePage(project)

    deadline = time.time() + 5.0
    while page._protocol_thread is not None and time.time() < deadline:
        qt_app.processEvents()
        time.sleep(0.01)

    assert not hasattr(page, "template_button")
    assert page.run_filter.currentData() == "ac52"
    assert page.plot.events
    assert all(row["source"] == "registered_protocol" for row in page.plot.events)
    assert not hasattr(page, "import_button")
    assert not hasattr(page, "add_button")
    assert page.approve_selected_run.text() == "Approve selected experiment"
    assert page.approve_all_runs.text() == "Approve all experiments and complete"
    page.close()
