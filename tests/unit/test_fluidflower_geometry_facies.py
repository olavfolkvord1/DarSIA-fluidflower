from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from PIL import Image
import pytest

from darsia.fluidflower.geometry_facies import (
    GeometryLabelEditSession,
    GeometryRevisionStore,
    create_optical_facies_draft,
    default_run_geometry_config,
    import_darsia_rig,
    load_geometry_preview,
    load_registered_geometry_image,
    resolve_run_geometry_artifacts,
    screen_layer_boundary_change,
    validate_geometry_config,
)
from darsia.fluidflower.project import FluidFlowerProject
from darsia.fluidflower.protocol_schema import write_protocol_workbook_template


def _save_field(path: Path, array: np.ndarray) -> None:
    metadata = {
        "space_dim": 2,
        "indexing": "ij",
        "dimensions": [0.49, 0.9],
        "origin": "[0.0, 0.49]",
    }
    np.savez_compressed(path, array=array, metadata=metadata)


def create_rig(tmp_path: Path) -> tuple[Path, Path, Path]:
    rig = tmp_path / "rig"
    rig.mkdir()
    baseline = np.zeros((6, 8, 3), dtype=np.uint8)
    baseline[..., 2] = 140
    labels = np.zeros((6, 8), dtype=np.int16)
    labels[2:, :4] = 1
    labels[2:, 4:] = 2
    facies = np.zeros_like(labels)
    facies[labels == 1] = 1
    facies[labels == 2] = 2
    depth = np.full(labels.shape, 0.01, dtype=np.float32)
    porosity = np.zeros(labels.shape, dtype=np.float32)
    porosity[labels == 1] = 0.40
    porosity[labels == 2] = 0.45
    permeability = np.zeros(labels.shape, dtype=np.float32)
    permeability[labels == 1] = 44.0
    permeability[labels == 2] = 473.0
    for name, array in (
        ("baseline.npz", baseline),
        ("labels.npz", labels),
        ("facies.npz", facies),
        ("depth.npz", depth),
        ("porosity.npz", porosity),
        ("permeability.npz", permeability),
    ):
        _save_field(rig / name, array)
    properties = tmp_path / "facies.csv"
    pd.DataFrame(
        [
            {"id": 0, "name": "water", "porosity": np.nan, "permeability": np.nan},
            {"id": 1, "name": "ESF", "porosity": 0.40, "permeability": 44.0},
            {"id": 2, "name": "C", "porosity": 0.45, "permeability": 473.0},
        ]
    ).to_csv(properties, index=False)
    protocol = tmp_path / "injection.csv"
    pd.DataFrame(
        [
            {"location_x": 0.30, "location_y": 0.10, "rate_sccm": 2.0},
            {"location_x": 0.65, "location_y": 0.25, "rate_sccm": 2.0},
        ]
    ).to_csv(protocol, index=False)
    return rig, properties, protocol


def create_project(tmp_path: Path) -> FluidFlowerProject:
    return FluidFlowerProject.create(
        tmp_path / "project",
        name="Geometry test",
        series_id="geometry-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )


def test_imported_darsia_rig_preserves_layers_and_protocol_injectors(
    tmp_path: Path,
) -> None:
    rig, properties, protocol = create_rig(tmp_path)
    config = import_darsia_rig(
        rig,
        facies_properties_path=properties,
        protocol_path=protocol,
        rig_type="medium",
    )

    assert config["width_m"] == 0.9
    assert config["height_m"] == 0.49
    assert [row["label_id"] for row in config["regions"]] == [0, 1, 2]
    assert [row["facies_id"] for row in config["regions"]] == [0, 1, 2]
    assert config["regions"][0]["analysis_role"] == "free_water"
    assert config["protocol_source_path"] == str(protocol.resolve())
    assert len(config["injectors"]) == 2
    assert config["injectors"][1]["x_m"] == 0.65


def test_excel_injection_points_can_be_positioned_later_in_gui(
    tmp_path: Path,
) -> None:
    rig, properties, _legacy_protocol = create_rig(tmp_path)
    protocol = write_protocol_workbook_template(
        tmp_path / "protocol.xlsx",
        layout="single_run",
        run_ids=("ac52",),
    )

    pending = import_darsia_rig(
        rig,
        facies_properties_path=properties,
        protocol_path=protocol,
        rig_type="medium",
    )

    assert [row["injector_id"] for row in pending["injectors"]] == ["ip1", "ip2"]
    assert all(
        row["x_m"] is None and row["y_m"] is None for row in pending["injectors"]
    )
    assert "invalid_injector_position" in {
        issue["code"] for issue in validate_geometry_config(pending)["issues"]
    }


def test_excel_position_is_visible_and_gui_override_is_reported(
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    rig, properties, _legacy_protocol = create_rig(tmp_path)
    protocol = write_protocol_workbook_template(
        tmp_path / "protocol.xlsx",
        layout="single_run",
        run_ids=("ac52",),
    )
    workbook = load_workbook(protocol)
    injection_points = workbook["Injection points"]
    headers = [cell.value for cell in injection_points[1]]
    x_column = headers.index("x_m") + 1
    y_column = headers.index("y_m") + 1
    injection_points.cell(2, x_column, 0.30)
    injection_points.cell(2, y_column, 0.10)
    injection_points.cell(3, x_column, 0.65)
    injection_points.cell(3, y_column, 0.25)
    workbook.save(protocol)

    config = import_darsia_rig(
        rig,
        facies_properties_path=properties,
        protocol_path=protocol,
        rig_type="medium",
    )

    assert config["injectors"][0]["injector_id"] == "ip1"
    assert config["injectors"][0]["x_m"] == 0.30
    assert config["injectors"][0]["protocol_x_m"] == 0.30
    config["injectors"][0]["x_m"] = 0.35
    codes = {issue["code"] for issue in validate_geometry_config(config)["issues"]}
    assert "injector_position_differs_from_protocol" in codes


def test_geometry_validation_calculates_physical_summary(tmp_path: Path) -> None:
    rig, properties, protocol = create_rig(tmp_path)
    config = import_darsia_rig(
        rig,
        facies_properties_path=properties,
        protocol_path=protocol,
        rig_type="medium",
    )
    report = validate_geometry_config(config)

    assert report["summary"]["error_count"] == 0
    assert report["summary"]["label_count"] == 3
    assert report["summary"]["facies_count"] == 3
    assert report["summary"]["injector_count"] == 2
    assert report["summary"]["active_area_fraction"] == 2.0 / 3.0
    assert report["summary"]["estimated_pore_volume_m3"] > 0.0
    assert report["summary"]["depth"]["median_m"] == pytest.approx(0.01)


def test_geometry_validation_rejects_shape_and_injector_errors(tmp_path: Path) -> None:
    rig, properties, protocol = create_rig(tmp_path)
    config = import_darsia_rig(
        rig,
        facies_properties_path=properties,
        protocol_path=protocol,
    )
    _save_field(rig / "depth.npz", np.full((5, 8), 0.01, dtype=np.float32))
    config["injectors"][0]["x_m"] = 1.5
    report = validate_geometry_config(config)
    codes = {issue["code"] for issue in report["issues"]}

    assert "shape_mismatch" in codes
    assert "injector_outside_rig" in codes
    assert report["summary"]["error_count"] >= 2


def test_shared_facies_requires_consistent_properties(tmp_path: Path) -> None:
    rig, properties, protocol = create_rig(tmp_path)
    config = import_darsia_rig(
        rig,
        facies_properties_path=properties,
        protocol_path=protocol,
    )
    config["regions"][2]["facies_id"] = 1
    report = validate_geometry_config(config)
    assert "inconsistent_shared_facies" in {issue["code"] for issue in report["issues"]}


def test_revision_store_materializes_canonical_facies(tmp_path: Path) -> None:
    rig, properties, protocol = create_rig(tmp_path)
    project = create_project(tmp_path)
    config = import_darsia_rig(
        rig,
        facies_properties_path=properties,
        protocol_path=protocol,
        rig_type="medium",
    )
    validation = validate_geometry_config(config)
    store = GeometryRevisionStore(project)
    revision = store.save_revision(
        config,
        actor="Reviewer",
        reason="Validated test geometry",
        validation=validation,
    )

    facies_path = Path(revision["artifacts"]["facies_path"])
    depth_path = Path(revision["artifacts"]["depth_path"])
    properties_path = Path(revision["artifacts"]["facies_properties_path"])
    assert facies_path.is_file()
    assert depth_path.is_file()
    assert properties_path.is_file()
    assert depth_path.parent == facies_path.parent
    assert revision["configuration"]["depth_path"] == str(depth_path)
    assert revision["configuration"]["depth_artifact_path"] == str(depth_path)
    assert store.current_config()["depth_path"] == str(depth_path)
    assert store.current()["revision_id"] == revision["revision_id"]
    assert project.manifest["rig"]["geometry_revision_id"] == revision["revision_id"]
    assert project.manifest["rig"]["depth"]["path"] == str(depth_path)
    assert revision["source_files"]["depth_path"]["path"] == str(
        (rig / "depth.npz").resolve()
    )
    assert revision["source_files"]["depth_path"]["sha256"]
    assert revision["artifacts"]["depth_file"]["sha256"] == revision[
        "source_files"
    ]["depth_materialization_input_path"]["sha256"]
    assert revision["source_files"]["protocol_source_path"]["path"] == str(
        protocol.resolve()
    )
    with np.load(facies_path, allow_pickle=False) as archive:
        assert set(np.unique(archive["array"]).tolist()) == {0, 1, 2}

    (rig / "depth.npz").unlink()
    preview = load_geometry_preview(store.current_config())
    assert preview["depth"].shape == (6, 8)
    assert np.allclose(preview["depth"], 0.01)


def test_new_geometry_revision_preserves_original_depth_provenance(
    tmp_path: Path,
) -> None:
    rig, properties, protocol = create_rig(tmp_path)
    project = create_project(tmp_path)
    config = import_darsia_rig(
        rig,
        facies_properties_path=properties,
        protocol_path=protocol,
        rig_type="medium",
    )
    store = GeometryRevisionStore(project)
    first = store.save_revision(
        config,
        actor="Reviewer",
        reason="Initial geometry",
    )
    original_reference = first["source_files"]["depth_path"]

    second_config = store.current_config()
    second_config["regions"][1]["name"] = "Renamed ESF"
    second = store.save_revision(
        second_config,
        actor="Reviewer",
        reason="Metadata-only geometry update",
    )

    assert second["source_files"]["depth_path"] == original_reference
    assert second["source_files"]["depth_materialization_input_path"][
        "path"
    ] == first["artifacts"]["depth_path"]
    assert Path(second["artifacts"]["depth_path"]).is_file()
    assert second["artifacts"]["depth_file"]["sha256"] == first["artifacts"][
        "depth_file"
    ]["sha256"]


def test_preview_is_downsampled_and_keeps_discrete_maps(tmp_path: Path) -> None:
    rig, properties, protocol = create_rig(tmp_path)
    config = import_darsia_rig(
        rig,
        facies_properties_path=properties,
        protocol_path=protocol,
    )
    preview = load_geometry_preview(config, max_width=4, max_height=4)

    assert preview["baseline"].shape == (3, 4, 3)
    assert preview["labels"].shape == (3, 4)
    assert preview["facies"].shape == (3, 4)
    assert set(np.unique(preview["labels"]).tolist()).issubset({0, 1, 2})


def test_run_geometry_materializes_layer_override_without_property_changes(
    tmp_path: Path,
) -> None:
    rig, properties, protocol = create_rig(tmp_path)
    project = create_project(tmp_path)
    config = import_darsia_rig(
        rig,
        facies_properties_path=properties,
        protocol_path=protocol,
        rig_type="medium",
    )
    source = tmp_path / "ac17-registered.png"
    Image.fromarray(np.full((6, 8, 3), (60, 100, 150), dtype=np.uint8)).save(source)
    labels = np.load(rig / "labels.npz")["array"].copy()
    labels[2, 3:5] = 2
    labels_path = tmp_path / "ac17-labels.npz"
    _save_field(labels_path, labels)
    run_geometry = default_run_geometry_config()
    run_geometry["runs"] = [
        {
            "run_id": "ac17",
            "enabled": True,
            "acquisition_state": "indicator",
            "source_image_path": str(source),
            "labels_path": str(labels_path),
            "facies_path": "",
            "property_change_proposals": [],
        }
    ]
    config["run_geometry"] = run_geometry

    report = validate_geometry_config(config)
    codes = {row["code"] for row in report["issues"]}
    assert report["summary"]["error_count"] == 0
    assert report["summary"]["run_geometry_count"] == 1
    assert "indicator_facies_carryover_risk" in codes

    revision = GeometryRevisionStore(project).save_revision(
        config,
        actor="Reviewer",
        reason="Run-specific compressed layers",
        validation=report,
    )
    resolved = resolve_run_geometry_artifacts(revision, "ac17")
    assert resolved["scope"] == "run_override"
    assert resolved["property_policy"] == "inherit_approved_series_properties"
    assert Path(resolved["labels_path"]).is_file()
    assert Path(resolved["facies_path"]).is_file()
    preview = load_geometry_preview(config, run_id="ac17")
    assert preview["run_id"] == "ac17"
    assert preview["labels"].shape == (6, 8)


def test_layer_change_screen_warns_without_selecting_geometry() -> None:
    reference = np.zeros((100, 160), dtype=np.int16)
    reference[20:50] = 1
    reference[50:80] = 2
    candidate = reference.copy()
    candidate[50:58] = 1

    report = screen_layer_boundary_change(
        reference,
        candidate,
        width_m=0.9,
        height_m=0.49,
        layer_label_ids={1, 2},
        warning_displacement_fraction_of_height=0.005,
        minimum_affected_boundary_fraction=0.10,
    )

    assert report["status"] == "substantial_change"
    assert report["substantial_change"] is True
    assert report["p95_displacement_m"] > report["threshold_m"]


def test_series_geometry_policy_keeps_run_map_advisory(tmp_path: Path) -> None:
    rig, properties, protocol = create_rig(tmp_path)
    project = create_project(tmp_path)
    config = import_darsia_rig(
        rig,
        facies_properties_path=properties,
        protocol_path=protocol,
        rig_type="medium",
    )
    run_labels = np.load(rig / "labels.npz")["array"].copy()
    run_labels[2:, 4:] = 1
    labels_path = tmp_path / "ac18-layer-review.npz"
    _save_field(labels_path, run_labels)
    run_geometry = default_run_geometry_config()
    run_geometry["geometry_policy"] = "series_reference_only"
    run_geometry["runs"] = [
        {
            "run_id": "ac18",
            "enabled": True,
            "acquisition_state": "clearwater",
            "geometry_choice": "run_override",
            "source_image_path": "",
            "labels_path": str(labels_path),
            "facies_path": "",
            "property_change_proposals": [],
        }
    ]
    config["run_geometry"] = run_geometry

    report = validate_geometry_config(config)
    codes = {row["code"] for row in report["issues"]}
    assert report["summary"]["error_count"] == 0
    assert report["summary"]["selected_run_override_count"] == 0
    assert "run_geometry_override_not_selected" in codes

    revision = GeometryRevisionStore(project).save_revision(
        config,
        actor="Reviewer",
        reason="Retain one shared series geometry",
        validation=report,
    )
    resolved = resolve_run_geometry_artifacts(revision, "ac18")
    assert resolved["scope"] == "series_reference"


def test_registered_review_frame_uses_same_geometry_grid(tmp_path: Path) -> None:
    rig, properties, protocol = create_rig(tmp_path)
    config = import_darsia_rig(
        rig,
        facies_properties_path=properties,
        protocol_path=protocol,
    )
    review = tmp_path / "review.png"
    Image.fromarray(np.full((6, 8, 3), (80, 120, 160), dtype=np.uint8)).save(review)
    config["display_image_path"] = str(review)

    image, provenance = load_registered_geometry_image(config, review)
    preview = load_geometry_preview(config)

    assert image.shape == (6, 8, 3)
    assert provenance["mode"] == "already_registered"
    assert preview["review_frame"].shape == preview["baseline"].shape


def test_optical_draft_is_review_required_and_reproducible(tmp_path: Path) -> None:
    source = np.zeros((120, 180, 3), dtype=np.uint8)
    source[:40] = (20, 50, 130)
    source[40:80] = (170, 150, 70)
    source[80:] = (70, 120, 80)
    source_path = tmp_path / "unknown-rig.png"
    Image.fromarray(source).save(source_path)

    config = create_optical_facies_draft(
        source_path,
        tmp_path / "draft",
        width_m=0.9,
        height_m=0.49,
        depth_m=0.01,
        optical_classes=3,
        rig_type="medium",
    )
    report = validate_geometry_config(config)
    preview = load_geometry_preview(config)

    assert config["facies_generation"]["status"] == "unreviewed_optical_draft"
    assert config["facies_generation"]["optical_classes"] == 3
    assert all(not row["enabled"] for row in config["regions"])
    assert all(row["analysis_role"] == "excluded" for row in config["regions"])
    assert set(np.unique(preview["facies"]).tolist()) == {1, 2, 3}
    assert report["summary"]["error_count"] == 0
    assert "no_active_porous_regions" in {issue["code"] for issue in report["issues"]}


def test_label_edits_are_reversible_full_resolution_and_materialized(
    tmp_path: Path,
) -> None:
    rig, properties, protocol = create_rig(tmp_path)
    config = import_darsia_rig(
        rig,
        facies_properties_path=properties,
        protocol_path=protocol,
        rig_type="medium",
    )
    source_labels = np.load(rig / "labels.npz")["array"].copy()
    session = GeometryLabelEditSession.create(config, tmp_path / "edit-sessions")

    polygon = session.apply(
        kind="polygon",
        points_m=[[0.0, 0.0], [0.42, 0.0], [0.42, 0.20], [0.0, 0.20]],
        target_label=2,
        actor="Reviewer",
    )
    after_polygon = session.labels.copy()
    brush = session.apply(
        kind="brush",
        points_m=[[0.55, 0.42], [0.82, 0.42]],
        target_label=1,
        actor="Reviewer",
        brush_radius_m=0.04,
    )
    after_brush = session.labels.copy()

    assert polygon["affected_pixel_count"] > 0
    assert brush["affected_pixel_count"] > 0
    assert not np.array_equal(after_polygon, source_labels)
    assert not np.array_equal(after_brush, after_polygon)
    assert np.array_equal(np.load(rig / "labels.npz")["array"], source_labels)

    assert session.undo(actor="Reviewer") is not None
    assert np.array_equal(session.labels, after_polygon)
    assert session.redo(actor="Reviewer") is not None
    assert np.array_equal(session.labels, after_brush)

    resumed_config = dict(config)
    resumed_config["label_edit_session"] = session.config_reference()
    resumed = GeometryLabelEditSession.create(
        resumed_config, tmp_path / "edit-sessions"
    )
    assert np.array_equal(resumed.labels, after_brush)
    assert resumed.active_operation_count() == 2

    materialization = resumed.materialize()
    with np.load(materialization["labels_path"], allow_pickle=False) as archive:
        assert np.array_equal(archive["array"], after_brush)
    assert np.array_equal(np.load(rig / "labels.npz")["array"], source_labels)
    assert [event["kind"] for event in resumed.events] == [
        "polygon",
        "brush",
        "undo",
        "redo",
    ]


def test_geometry_revision_records_label_edit_provenance(tmp_path: Path) -> None:
    rig, properties, protocol = create_rig(tmp_path)
    project = create_project(tmp_path)
    config = import_darsia_rig(
        rig,
        facies_properties_path=properties,
        protocol_path=protocol,
        rig_type="medium",
    )
    session = GeometryLabelEditSession.create(config, tmp_path / "edit-sessions")
    session.apply(
        kind="brush",
        points_m=[[0.20, 0.10]],
        target_label=2,
        actor="Reviewer",
        brush_radius_m=0.04,
    )
    materialization = session.materialize()
    reference = session.config_reference()
    reference["materialization_manifest_path"] = materialization["manifest_path"]
    reference["materialized_active_operation_count"] = 1
    config["labels_path"] = materialization["labels_path"]
    config["facies_source_path"] = ""
    config["label_edit_session"] = reference

    validation = validate_geometry_config(config)
    revision = GeometryRevisionStore(project).save_revision(
        config,
        actor="Reviewer",
        reason="Edited one physical boundary",
        validation=validation,
    )

    assert revision["source_files"]["label_edit_base_labels_path"]["path"] == str(
        (rig / "labels.npz").resolve()
    )
    assert revision["source_files"]["label_edit_events_path"]["path"] == str(
        session.events_path.resolve()
    )
    assert revision["source_files"]["label_edit_materialization_manifest_path"][
        "path"
    ] == str(Path(materialization["manifest_path"]).resolve())
