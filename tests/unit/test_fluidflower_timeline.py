from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from darsia.fluidflower.project import FluidFlowerProject
from darsia.fluidflower.protocol_schema import (
    read_protocol_tables,
    write_protocol_file_set_template,
    write_protocol_workbook_template,
)
from darsia.fluidflower.timeline import TimelineStore


def create_project(tmp_path: Path) -> FluidFlowerProject:
    return FluidFlowerProject.create(
        tmp_path / "project",
        name="Timeline test",
        series_id="timeline-test",
        template_id="ff_ac",
        rig_type="medium",
        created_by="Reviewer",
    )


def test_timeline_overrides_do_not_modify_raw_event(tmp_path: Path) -> None:
    project = create_project(tmp_path)
    timeline = TimelineStore(project)
    event = timeline.add_event(
        run_id="ac60",
        source="protocol",
        event_type="injection_start",
        time_h=0.0,
        label="Injection 1",
        actor="Reviewer",
    )

    timeline.set_override(
        event["event_id"],
        {"time_h": 0.25, "note": "Camera clock offset"},
        actor="Reviewer",
        reason="Aligned against the first injection image",
    )

    raw = timeline.raw_events()[0]
    effective = timeline.effective_events()[0]
    assert raw["raw_time_h"] == 0.0
    assert "time_h" not in raw
    assert effective["time_h"] == 0.25
    assert effective["note"] == "Camera clock offset"


def test_timeline_clear_restores_raw_value_without_erasing_history(
    tmp_path: Path,
) -> None:
    timeline = TimelineStore(create_project(tmp_path))
    event = timeline.add_event(
        run_id="run18",
        source="images",
        event_type="image_capture",
        time_h=1.0,
        actor="Reviewer",
    )
    timeline.set_override(
        event["event_id"],
        {"time_h": 1.5},
        actor="Reviewer",
        reason="First alignment",
    )
    timeline.clear_overrides(
        event["event_id"],
        ["time_h"],
        actor="Reviewer",
        reason="Revert local alignment",
    )

    assert timeline.effective_events()[0]["time_h"] == 1.0
    assert len(timeline.overrides()) == 2


def test_protocol_interval_import_is_idempotent(tmp_path: Path) -> None:
    project = create_project(tmp_path)
    source = tmp_path / "injection.csv"
    source.write_text(
        "id,start,end,rate_ml_min,label\n"
        "1,2026-01-01 12:00:00,2026-01-01 12:45:00,2.0,I1\n",
        encoding="utf-8",
    )
    timeline = TimelineStore(project)
    first = timeline.import_table(
        source,
        source="injection",
        run_id="ac60",
        actor="Reviewer",
    )
    second = timeline.import_table(
        source,
        source="injection",
        run_id="ac60",
        actor="Reviewer",
    )

    assert first == {"import_id": first["import_id"], "added": 2, "skipped": 0}
    assert second["added"] == 0
    assert second["skipped"] == 2
    events = timeline.effective_events()
    assert [event["event_type"] for event in events] == [
        "injection_start",
        "injection_end",
    ]
    assert events[1]["time_h"] == 0.75


def test_timeline_qc_reports_missing_and_duplicate_times(tmp_path: Path) -> None:
    timeline = TimelineStore(create_project(tmp_path))
    for time_h in (None, 2.0, 2.0):
        timeline.add_event(
            run_id="ac60",
            source="images",
            event_type="image_capture",
            time_h=time_h,
            actor="Reviewer",
        )
    issues = timeline.qc_issues()
    assert sum(issue["code"] == "missing_time" for issue in issues) == 1
    assert sum(issue["code"] == "duplicate_time" for issue in issues) == 2


def test_protocol_workbook_imports_all_run_sheets(tmp_path: Path) -> None:
    workbook = write_protocol_workbook_template(
        tmp_path / "protocol.xlsx",
        layout="per_run_sheets",
        run_ids=("run18", "run20"),
    )
    timeline = TimelineStore(create_project(tmp_path))

    result = timeline.import_table(
        workbook,
        source="protocol",
        actor="Reviewer",
    )

    assert result["added"] == 28
    events = timeline.effective_events()
    assert {event["run_id"] for event in events} == {"run18", "run20"}
    assert {
        str((event.get("metadata") or {}).get("source_sheet") or "").casefold()
        for event in events
    } == {"run18", "run20"}


def test_registered_protocols_are_imported_once_from_project_mapping(
    tmp_path: Path,
) -> None:
    workbook = write_protocol_workbook_template(
        tmp_path / "ac52_protocol.xlsx",
        layout="single_run",
        run_ids=("ac52",),
    )
    project = create_project(tmp_path)
    project.manifest["runs"] = [
        {
            "run_id": "ac52",
            "source_mapping": {
                "protocol_table": str(workbook),
                "protocol_tables": [str(workbook)],
            },
        }
    ]
    project._write_manifest()
    timeline = TimelineStore(project)

    first = timeline.import_registered_protocols(actor="Reviewer")
    second = timeline.import_registered_protocols(actor="Reviewer")

    assert first["registered_source_count"] == 1
    assert first["added"] > 0
    assert first["errors"] == []
    assert second["added"] == 0
    assert second["skipped"] == first["added"]
    assert {row["source"] for row in timeline.effective_events()} == {
        "registered_protocol"
    }
    assert {row["run_id"] for row in timeline.effective_events()} == {"ac52"}


def test_combined_protocol_template_contains_metadata_and_event_guidance(
    tmp_path: Path,
) -> None:
    workbook_path = write_protocol_workbook_template(
        tmp_path / "combined.xlsx",
        layout="combined",
        run_ids=("ac52", "ac60"),
    )

    workbook = load_workbook(workbook_path, data_only=True)

    assert workbook.sheetnames == [
        "README",
        "Runs",
        "Injection points",
        "Event types",
        "Events",
    ]
    run_headers = [cell.value for cell in next(workbook["Runs"].iter_rows())]
    event_headers = [cell.value for cell in next(workbook["Events"].iter_rows())]
    assert "experiment_date" in run_headers
    assert "timezone" in run_headers
    assert "indicator_initial_ph" in run_headers
    assert "start_datetime" in event_headers
    assert "first_image" in event_headers
    assert "measurement_location" in event_headers

    injection_rows = list(workbook["Injection points"].iter_rows(values_only=True))
    assert len(injection_rows) == 5
    assert {row[0] for row in injection_rows[1:]} == {"ac52", "ac60"}
    injection_headers = [cell.value for cell in workbook["Injection points"][1]]
    assert "z_m" not in injection_headers
    assert {row[1] for row in injection_rows[1:]} == {"ip1", "ip2"}
    assert all(row[3] is None and row[4] is None for row in injection_rows[1:])
    x_header = workbook["Injection points"][1][injection_headers.index("x_m")]
    assert x_header.comment is not None
    assert "Leave blank" in x_header.comment.text
    readme_text = "\n".join(
        str(cell.value or "") for row in workbook["README"] for cell in row
    )
    assert "x_m and y_m are optional" in readme_text


def test_protocol_folder_infers_run_from_filename(tmp_path: Path) -> None:
    folder = tmp_path / "protocols"
    folder.mkdir()
    for run_id in ("ac17", "ac18"):
        (folder / f"{run_id}.csv").write_text(
            "event_type,time_h,label\nimage_capture,0.0,first\n",
            encoding="utf-8",
        )
    timeline = TimelineStore(create_project(tmp_path))

    result = timeline.import_folder(folder, source="protocol", actor="Reviewer")

    assert result["added"] == 2
    assert {event["run_id"] for event in timeline.effective_events()} == {
        "ac17",
        "ac18",
    }


def test_per_run_file_template_matches_folder_organization(tmp_path: Path) -> None:
    folder = write_protocol_file_set_template(
        tmp_path / "protocol-example",
        run_ids=("ac17", "ac18"),
    )

    assert {path.name for path in folder.glob("*.xlsx")} == {
        "ac17.xlsx",
        "ac18.xlsx",
    }
    rows = read_protocol_tables(folder)
    assert len(rows) == 20
    assert {row["run_id"] for row in rows} == {"ac17", "ac18"}


def test_protocol_v2_resolves_clock_times_and_multiple_days(tmp_path: Path) -> None:
    workbook = write_protocol_workbook_template(
        tmp_path / "protocol-v2.xlsx",
        layout="single_run",
        run_ids=("run18",),
    )
    timeline = TimelineStore(create_project(tmp_path))

    result = timeline.import_table(
        workbook,
        source="protocol",
        actor="Reviewer",
    )

    assert result["added"] == 14
    events = timeline.effective_events()
    second_injection = [
        event
        for event in events
        if event["event_type"] == "injection_steady_start" and event["time_h"] == 24.0
    ]
    assert len(second_injection) == 1
    assert second_injection[0]["unit"] == "ml/min"
    assert second_injection[0]["value"] == 5.0
    assert second_injection[0]["raw_timestamp"].startswith("2026-01-16T09:00:00")

    ramp = [
        event
        for event in events
        if event["event_type"]
        in {
            "injection_ramp_up_start",
            "injection_ramp_up_end",
        }
    ]
    assert [event["value"] for event in ramp] == [0.1, 2.0]
    assert {event["unit"] for event in ramp} == {"ml/min"}

    ph_events = [event for event in events if event["event_type"] == "ph_measurement"]
    assert len(ph_events) == 2
    assert {event["unit"] for event in ph_events} == {"pH"}
    assert {
        str((event["metadata"]["raw_row"]).get("measurement_location") or "")
        for event in ph_events
    } == {"indicator reservoir", "rig before injection"}


def test_protocol_v2_uses_first_image_for_camera_interval(tmp_path: Path) -> None:
    source = tmp_path / "camera.csv"
    source.write_text(
        "run_id,event_type,experiment_date,datetime,time_h,value,unit,first_image\n"
        "ac52,experiment_start,2026-01-15,09:00:00,0,,,\n"
        "ac52,camera_interval_start,2026-01-15,10:00:00,1,300,s,DSC00121.JPG\n",
        encoding="utf-8",
    )
    timeline = TimelineStore(create_project(tmp_path))

    timeline.import_table(source, source="protocol", actor="Reviewer")

    camera_event = next(
        event
        for event in timeline.effective_events()
        if event["event_type"] == "camera_interval_start"
    )
    assert camera_event["path"] == "DSC00121.JPG"
    assert camera_event["value"] == 300.0
    assert camera_event["unit"] == "s"


def test_protocol_qc_reports_relative_and_dated_time_disagreement(
    tmp_path: Path,
) -> None:
    source = tmp_path / "mismatch.csv"
    source.write_text(
        "run_id,event_type,experiment_date,datetime,time_h\n"
        "ac52,experiment_start,2026-01-15,09:00:00,0\n"
        "ac52,intervention,2026-01-15,10:00:00,4\n",
        encoding="utf-8",
    )
    timeline = TimelineStore(create_project(tmp_path))

    timeline.import_table(source, source="protocol", actor="Reviewer")

    events = timeline.effective_events()
    intervention = next(
        event for event in events if event["event_type"] == "intervention"
    )
    assert intervention["time_h"] == 1.0
    assert any(
        issue["event_id"] == intervention["event_id"]
        and issue["code"] == "time_source_disagreement"
        for issue in timeline.qc_issues()
    )
