"""Read iWeight KINETICS exports from FluidFlower overflow balances.

The iWeight balance software writes one worksheet per logging session::

    2026/05/05   14:04:08            <- export time on the balance PC
    Riga 1 ... Riga 4                <- free header lines
    Weighing Mode:  KINETICS
    Interval:       1       s
    Max Time:       604800  s
    WEIGHTS
    NUM  WEIGHT  UM  TIME            <- TIME is a clock time without a date
    1    59.5    g   12:45:10
    ...
    Min  2217.2                      <- optional summary rows

Sample times carry no date. The date of every sample is reconstructed from the
export time in the header and from midnight roll-overs in the TIME column.
When logging stopped long before the export (for example because ``Max Time``
was reached), the number of whole days between the last sample and the export
cannot be read from the file. :func:`anchor_session` therefore takes an
explicit ``day_shift``, and :func:`rank_day_shifts` scores candidate shifts
against known events such as injection starts.

The module only needs the standard library, numpy, pandas and openpyxl. It does
not import darsia, so it also runs as a stand-alone script on Python 3.10::

    python iweight_export.py inspect export1.xlsx export2.xlsx
"""

from __future__ import annotations

import argparse
import dataclasses
import datetime as dt
import json
import re
import sys
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

import numpy as np
import pandas as pd

FORMAT = "darsia-fluidflower-iweight-kinetics-v1"

_EXPORT_TIME = re.compile(
    r"(\d{4})[/.-](\d{1,2})[/.-](\d{1,2})\s+(\d{1,2}):(\d{2}):(\d{2})"
)
_CLOCK = re.compile(r"^\s*(\d{1,2}):(\d{2}):(\d{2}(?:\.\d+)?)\s*$")
_UNIT_TO_G = {"g": 1.0, "kg": 1000.0, "mg": 0.001}
_DURATION_UNIT_TO_S = {"s": 1.0, "sec": 1.0, "ms": 0.001, "min": 60.0, "h": 3600.0}
_DAY_S = 86400.0
# A backward step of the clock larger than this marks a new day.
_ROLLOVER_S = 43200.0


class KineticsFormatError(ValueError):
    """Raised for worksheets that look like iWeight exports but cannot be used."""


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _is_number(value: Any) -> bool:
    if isinstance(value, bool) or value is None:
        return False
    if isinstance(value, (int, float, np.integer, np.floating)):
        return bool(np.isfinite(float(value)))
    try:
        return bool(np.isfinite(float(str(value).strip().replace(",", "."))))
    except ValueError:
        return False


def _number(value: Any) -> float:
    if isinstance(value, (int, float, np.integer, np.floating)):
        return float(value)
    return float(str(value).strip().replace(",", "."))


def parse_export_time(value: Any) -> dt.datetime | None:
    """Return the export time in an iWeight header cell, if there is one."""
    if isinstance(value, dt.datetime):
        return value
    match = _EXPORT_TIME.search(_text(value))
    if match is None:
        return None
    year, month, day, hour, minute, second = map(int, match.groups())
    return dt.datetime(year, month, day, hour, minute, second)


def clock_seconds(value: Any) -> float | None:
    """Seconds after midnight for a TIME cell, or ``None`` if it is not a time."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, dt.datetime):
        value = value.time()
    if isinstance(value, dt.time):
        return (
            value.hour * 3600.0
            + value.minute * 60.0
            + value.second
            + value.microsecond / 1e6
        )
    if isinstance(value, dt.timedelta):
        seconds = value.total_seconds()
        return seconds if 0.0 <= seconds < _DAY_S else None
    if isinstance(value, (int, float, np.integer, np.floating)):
        fraction = float(value)  # Excel stores clock times as fractions of a day
        return round(fraction * _DAY_S, 3) if 0.0 <= fraction < 1.0 else None
    match = _CLOCK.match(str(value))
    if match is None:
        return None
    hour, minute, second = int(match[1]), int(match[2]), float(match[3])
    if hour >= 24 or minute >= 60 or second >= 60.0:
        return None
    return hour * 3600.0 + minute * 60.0 + second


def _duration_s(values: Sequence[str]) -> float | None:
    if not values or not _is_number(values[0]):
        return None
    unit = values[1].casefold() if len(values) > 1 else "s"
    if unit not in _DURATION_UNIT_TO_S:
        raise KineticsFormatError(f"Unknown duration unit {values[1]!r}")
    return _number(values[0]) * _DURATION_UNIT_TO_S[unit]


def _grams_per_unit(unit: Any) -> float:
    key = _text(unit).casefold() or "g"
    if key not in _UNIT_TO_G:
        raise KineticsFormatError(f"Unsupported weight unit {unit!r}")
    return _UNIT_TO_G[key]


@dataclasses.dataclass(frozen=True, eq=False)
class KineticsSession:
    """One iWeight logging session (one worksheet).

    ``samples`` has the columns ``num``, ``weight_g``, ``clock_s`` (seconds
    after midnight), ``day`` (midnight roll-overs since the first sample) and
    ``session_s`` (seconds since the first sample).
    """

    source_path: str
    sheet: str
    export_time: dt.datetime
    weighing_mode: str
    interval_s: float | None
    max_time_s: float | None
    header_lines: tuple[str, ...]
    samples: pd.DataFrame
    ignored_rows: int
    summary: Mapping[str, float]
    clock_backsteps: int

    @property
    def n_samples(self) -> int:
        return int(len(self.samples))

    @property
    def duration_s(self) -> float:
        return float(self.samples["session_s"].iloc[-1])

    @property
    def max_sample_gap_s(self) -> float:
        steps = np.diff(self.samples["session_s"].to_numpy())
        return float(steps.max()) if steps.size else 0.0

    @property
    def num_clock_drift_s(self) -> float:
        """Largest difference between clock time and ``NUM x Interval``."""
        if not self.interval_s:
            return float("nan")
        num = self.samples["num"].to_numpy()
        expected = (num - num[0]) * self.interval_s
        return float(np.max(np.abs(self.samples["session_s"].to_numpy() - expected)))

    @property
    def reached_max_time(self) -> bool:
        """Whether logging ran until ``Max Time`` and then stopped by itself."""
        if not self.max_time_s:
            return False
        step = self.interval_s or 1.0
        return self.duration_s >= self.max_time_s - 2.0 * step

    def end_time(self, day_shift: int = 0, *, tolerance_s: float = 5.0) -> dt.datetime:
        """Time of the last sample.

        With ``day_shift=0`` the last sample is placed on the latest date that
        is not after the export time. Each unit of ``day_shift`` moves the
        whole session one day earlier.
        """
        if day_shift < 0:
            raise ValueError("day_shift must be zero or positive")
        last_clock = float(self.samples["clock_s"].iloc[-1])
        midnight = dt.datetime.combine(self.export_time.date(), dt.time())
        end = midnight + dt.timedelta(seconds=last_clock)
        if end > self.export_time + dt.timedelta(seconds=tolerance_s):
            end -= dt.timedelta(days=1)
        return end - dt.timedelta(days=int(day_shift))

    def start_time(self, day_shift: int = 0) -> dt.datetime:
        return self.end_time(day_shift) - dt.timedelta(seconds=self.duration_s)

    def export_gap_s(self, day_shift: int = 0) -> float:
        """Seconds between the last sample and the export."""
        return (self.export_time - self.end_time(day_shift)).total_seconds()

    def day_shift_is_certain(self, *, max_export_gap_s: float = 3600.0) -> bool:
        """True when the default ``day_shift=0`` needs no outside evidence.

        This is the case when logging was still running shortly before the
        export. If logging stopped at ``Max Time`` or the export came much
        later, whole days may have passed in between.
        """
        return (not self.reached_max_time) and self.export_gap_s(0) <= max_export_gap_s

    def describe(self, day_shift: int = 0) -> dict[str, Any]:
        first = self.samples.iloc[0]
        last = self.samples.iloc[-1]
        return {
            "format": FORMAT,
            "source_path": self.source_path,
            "sheet": self.sheet,
            "export_time": self.export_time.isoformat(sep=" "),
            "weighing_mode": self.weighing_mode,
            "interval_s": self.interval_s,
            "max_time_s": self.max_time_s,
            "n_samples": self.n_samples,
            "first_clock": _format_clock(float(first["clock_s"])),
            "last_clock": _format_clock(float(last["clock_s"])),
            "midnight_rollovers": int(last["day"]),
            "duration_s": self.duration_s,
            "max_sample_gap_s": self.max_sample_gap_s,
            "num_clock_drift_s": self.num_clock_drift_s,
            "clock_backsteps": self.clock_backsteps,
            "reached_max_time": self.reached_max_time,
            "day_shift": int(day_shift),
            "start_time": self.start_time(day_shift).isoformat(sep=" "),
            "end_time": self.end_time(day_shift).isoformat(sep=" "),
            "export_gap_s": self.export_gap_s(day_shift),
            "day_shift_certain": self.day_shift_is_certain(),
            "first_weight_g": float(first["weight_g"]),
            "last_weight_g": float(last["weight_g"]),
            "ignored_rows": self.ignored_rows,
            "summary": dict(self.summary),
        }


def _format_clock(seconds: float) -> str:
    seconds = int(round(seconds)) % int(_DAY_S)
    return f"{seconds // 3600:02d}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}"


def parse_kinetics_rows(
    rows: Iterable[Sequence[Any] | None],
    *,
    source_path: str = "",
    sheet: str = "",
) -> KineticsSession | None:
    """Parse the rows of one worksheet.

    Returns ``None`` when the rows are not an iWeight export (for example a
    sheet that someone added by hand). Rows below the column header that lack
    NUM, WEIGHT or TIME are counted in ``ignored_rows``, except summary rows
    such as ``Min``/``Max``, which are collected in ``summary``.
    """
    iterator = iter(rows)
    export_time: dt.datetime | None = None
    mode: str | None = None
    interval_s: float | None = None
    max_time_s: float | None = None
    header_lines: list[str] = []
    columns: dict[str, int] | None = None

    for row in iterator:
        cells = list(row) if row is not None else []
        texts = [_text(cell) for cell in cells]
        filled = [text for text in texts if text]
        if not filled:
            continue
        key = filled[0].rstrip(":").strip().casefold()
        if export_time is None and mode is None:
            found = next(
                (parse_export_time(c) for c in cells if parse_export_time(c)), None
            )
            if found is not None:
                export_time = found
                continue
        if key == "weighing mode":
            mode = filled[1] if len(filled) > 1 else ""
            continue
        if key == "interval":
            interval_s = _duration_s(filled[1:])
            continue
        if key == "max time":
            max_time_s = _duration_s(filled[1:])
            continue
        if key == "weights":
            continue
        upper = [text.upper() for text in texts]
        if "NUM" in upper and "WEIGHT" in upper and "TIME" in upper:
            columns = {
                name: upper.index(name)
                for name in ("NUM", "WEIGHT", "UM", "TIME")
                if name in upper
            }
            break
        if mode is None:
            header_lines.append(filled[0])

    if columns is None or mode is None:
        return None
    if export_time is None:
        raise KineticsFormatError(f"No export time in {source_path} [{sheet}]")
    if mode.upper() != "KINETICS":
        raise KineticsFormatError(
            f"Weighing mode {mode!r} is not supported in {source_path} [{sheet}]"
        )

    def cell(cells: Sequence[Any], name: str) -> Any:
        index = columns.get(name)
        return cells[index] if index is not None and index < len(cells) else None

    nums: list[int] = []
    weights: list[float] = []
    clocks: list[float] = []
    ignored = 0
    summary: dict[str, float] = {}
    for row in iterator:
        cells = list(row) if row is not None else []
        num, weight = cell(cells, "NUM"), cell(cells, "WEIGHT")
        clock = clock_seconds(cell(cells, "TIME"))
        if _is_number(num) and _is_number(weight) and clock is not None:
            nums.append(int(_number(num)))
            weights.append(_number(weight) * _grams_per_unit(cell(cells, "UM")))
            clocks.append(clock)
            continue
        filled = [cell_value for cell_value in cells if _text(cell_value)]
        if not filled:
            continue
        label = _text(filled[0]).rstrip(":")
        values = [value for value in filled[1:] if _is_number(value)]
        if label and not _is_number(label) and values:
            summary[label] = _number(values[0])
        else:
            ignored += 1

    if not nums:
        raise KineticsFormatError(f"No samples in {source_path} [{sheet}]")

    clock_array = np.asarray(clocks, dtype=float)
    steps = np.diff(clock_array)
    rollover = steps < -_ROLLOVER_S
    backsteps = int(np.count_nonzero((steps < 0.0) & ~rollover))
    day = np.concatenate([[0], np.cumsum(rollover)]).astype(np.int64)
    session_s = clock_array + day * _DAY_S - clock_array[0]
    samples = pd.DataFrame(
        {
            "num": np.asarray(nums, dtype=np.int64),
            "weight_g": np.asarray(weights, dtype=float),
            "clock_s": clock_array,
            "day": day,
            "session_s": session_s,
        }
    )
    return KineticsSession(
        source_path=str(source_path),
        sheet=str(sheet),
        export_time=export_time,
        weighing_mode=mode,
        interval_s=interval_s,
        max_time_s=max_time_s,
        header_lines=tuple(header_lines),
        samples=samples,
        ignored_rows=ignored,
        summary=summary,
        clock_backsteps=backsteps,
    )


def read_kinetics_workbook(
    path: str | Path, sheets: Sequence[str] | None = None
) -> list[KineticsSession]:
    """Read all iWeight sessions in an Excel workbook.

    Worksheets that are not iWeight exports are skipped. ``sheets`` limits the
    reading to the named worksheets.
    """
    from openpyxl import load_workbook

    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sessions = []
        for worksheet in workbook.worksheets:
            if sheets is not None and worksheet.title not in sheets:
                continue
            # iWeight writes no reliable dimension record; without this reset
            # openpyxl's read-only mode may see a single row.
            worksheet.reset_dimensions()
            session = parse_kinetics_rows(
                worksheet.iter_rows(values_only=True),
                source_path=str(path),
                sheet=worksheet.title,
            )
            if session is not None:
                sessions.append(session)
        return sessions
    finally:
        workbook.close()


@dataclasses.dataclass(frozen=True, eq=False)
class AnchoredSession:
    """A session placed on the balance-PC clock with a chosen ``day_shift``."""

    session: KineticsSession
    day_shift: int
    label: str
    frame: pd.DataFrame  # datetime, weight_g, num

    @property
    def start(self) -> pd.Timestamp:
        return self.frame["datetime"].iloc[0]

    @property
    def end(self) -> pd.Timestamp:
        return self.frame["datetime"].iloc[-1]


def anchor_session(
    session: KineticsSession, *, day_shift: int = 0, label: str | None = None
) -> AnchoredSession:
    """Give every sample of ``session`` an absolute time."""
    end = pd.Timestamp(session.end_time(day_shift))
    session_s = session.samples["session_s"].to_numpy()
    offsets = pd.to_timedelta(session_s - session_s[-1], unit="s")
    times = (end + offsets).round("ms")
    frame = pd.DataFrame(
        {
            "datetime": np.asarray(times, dtype="datetime64[ns]"),
            "weight_g": session.samples["weight_g"].to_numpy(),
            "num": session.samples["num"].to_numpy(),
        }
    )
    if label is None:
        label = f"{Path(session.source_path).name}[{session.sheet}]"
    return AnchoredSession(
        session=session, day_shift=int(day_shift), label=label, frame=frame
    )


def _empty_rises() -> pd.DataFrame:
    return pd.DataFrame(
        columns=[
            "onset_s",
            "crossing_s",
            "end_s",
            "peak_s",
            "duration_s",
            "baseline_g",
            "peak_g",
            "gain_g",
            "rate_g_s",
            "initial_rate_g_s",
            "baseline_n",
        ]
    )


def detect_rises(
    seconds: Sequence[float] | np.ndarray,
    weights: Sequence[float] | np.ndarray,
    *,
    window_s: float = 60.0,
    min_rate_g_s: float = 0.04,
    min_duration_s: float = 300.0,
    merge_gap_s: float = 60.0,
    baseline_s: float = 600.0,
    threshold_g: float = 0.5,
    fit_s: float = 120.0,
) -> pd.DataFrame:
    """Find sustained weight increases, such as displacement during injection.

    A sample is *rising* when the weight increased by at least
    ``min_rate_g_s`` per second over the trailing ``window_s``. Rising samples
    closer than ``merge_gap_s`` form one rise, which is kept when it lasts at
    least ``min_duration_s``. The onset is where a straight line fitted to the
    first ``fit_s`` after the weight crossed ``baseline + threshold_g`` meets the
    baseline (median weight over ``baseline_s`` before the rise).
    ``baseline_n`` is the number of samples behind that median; a rise that
    starts with the record has none and its onset is not reliable.

    Times are seconds on any monotonic axis; the result uses the same axis.
    """
    t = np.asarray(seconds, dtype=float)
    w = np.asarray(weights, dtype=float)
    if t.shape != w.shape:
        raise ValueError("seconds and weights must have the same length")
    if t.size < 3:
        return _empty_rises()
    if np.any(np.diff(t) < 0):
        raise ValueError("seconds must be sorted")
    j = np.searchsorted(t, t - window_s, side="left")
    span = t - t[j]
    valid = span >= 0.5 * window_s
    rate = np.full(t.shape, np.nan)
    rate[valid] = (w[valid] - w[j[valid]]) / span[valid]
    rising = np.flatnonzero(valid & (rate >= min_rate_g_s))
    if rising.size == 0:
        return _empty_rises()
    breaks = np.flatnonzero(np.diff(t[rising]) > merge_gap_s)
    starts = np.concatenate([[rising[0]], rising[breaks + 1]])
    ends = np.concatenate([rising[breaks], [rising[-1]]])

    rows = []
    for first, last in zip(starts, ends):
        window_start = t[first] - window_s
        if t[last] - window_start < min_duration_s:
            continue
        lo = np.searchsorted(t, window_start - baseline_s, side="left")
        hi = np.searchsorted(t, window_start, side="right")
        baseline = (
            float(np.median(w[lo:hi])) if hi - lo >= 5 else float(w[max(hi - 1, 0)])
        )
        k0 = np.searchsorted(t, window_start, side="left")
        above = np.flatnonzero(w[k0 : last + 1] > baseline + threshold_g)
        if above.size == 0:
            continue
        crossing = k0 + int(above[0])
        fit_end = np.searchsorted(t, t[crossing] + fit_s, side="right")
        fit_t = t[crossing:fit_end] - t[crossing]
        fit_w = w[crossing:fit_end]
        initial_rate = float("nan")
        onset = float(t[crossing])
        if fit_t.size >= 5 and np.ptp(fit_t) > 0:
            slope, intercept = np.polyfit(fit_t, fit_w, 1)
            initial_rate = float(slope)
            if slope > 0:
                estimate = t[crossing] + (baseline - intercept) / slope
                onset = float(min(max(estimate, t[crossing] - fit_s), t[crossing]))
        peak_end = np.searchsorted(t, t[last] + window_s, side="right")
        peak_index = crossing + int(np.argmax(w[crossing:peak_end]))
        rows.append(
            {
                "onset_s": onset,
                "crossing_s": float(t[crossing]),
                "end_s": float(t[last]),
                "peak_s": float(t[peak_index]),
                "duration_s": float(t[peak_index] - onset),
                "baseline_g": baseline,
                "peak_g": float(w[peak_index]),
                "gain_g": float(w[peak_index] - baseline),
                "rate_g_s": float(np.nanmedian(rate[first : last + 1])),
                "initial_rate_g_s": initial_rate,
                "baseline_n": int(hi - lo),
            }
        )
    return pd.DataFrame(rows) if rows else _empty_rises()


def detect_steps(
    seconds: Sequence[float] | np.ndarray,
    weights: Sequence[float] | np.ndarray,
    *,
    threshold_g: float = 20.0,
    span_s: float = 5.0,
    merge_s: float = 120.0,
    settle_s: float = 30.0,
) -> pd.DataFrame:
    """Find abrupt weight changes, such as an emptied or swapped container.

    A step is a change of at least ``threshold_g`` within ``span_s``. Steps
    closer than ``merge_s`` are merged into one event. ``before_g`` and
    ``after_g`` are medians over ``settle_s`` outside the event; ``delta_g`` is
    the net level change that an offset correction would remove.
    """
    columns = ["start_s", "end_s", "before_g", "after_g", "delta_g", "min_g", "max_g"]
    t = np.asarray(seconds, dtype=float)
    w = np.asarray(weights, dtype=float)
    if t.size < 2:
        return pd.DataFrame(columns=columns)
    j = np.searchsorted(t, t - span_s, side="left")
    flagged = np.flatnonzero(np.abs(w - w[j]) >= threshold_g)
    if flagged.size == 0:
        return pd.DataFrame(columns=columns)
    breaks = np.flatnonzero(np.diff(t[flagged]) > merge_s)
    starts = np.concatenate([[flagged[0]], flagged[breaks + 1]])
    ends = np.concatenate([flagged[breaks], [flagged[-1]]])
    rows = []
    for first, last in zip(starts, ends):
        a, b = int(j[first]), int(last)
        before = w[np.searchsorted(t, t[a] - settle_s, side="left") : a + 1]
        after = w[b : np.searchsorted(t, t[b] + settle_s, side="right")]
        before_g = float(np.median(before))
        after_g = float(np.median(after))
        rows.append(
            {
                "start_s": float(t[a]),
                "end_s": float(t[b]),
                "before_g": before_g,
                "after_g": after_g,
                "delta_g": after_g - before_g,
                "min_g": float(np.min(w[a : b + 1])),
                "max_g": float(np.max(w[a : b + 1])),
            }
        )
    return pd.DataFrame(rows, columns=columns)


def rank_day_shifts(
    session: KineticsSession,
    event_times: Iterable[Any],
    *,
    max_shift_days: int = 10,
    tolerance_s: float = 1800.0,
    **rise_options: Any,
) -> pd.DataFrame:
    """Score candidate ``day_shift`` values against known event times.

    Rises in the session (see :func:`detect_rises`) are matched to the
    nearest event, for instance injection starts read from the photographs.
    The best shift matches most rises, leaves fewest events inside the session
    without a rise, and has the smallest median time offset. The first row is
    the best candidate.
    """
    events = sorted(pd.Timestamp(value).to_pydatetime() for value in event_times)
    samples = session.samples
    rises = detect_rises(samples["session_s"], samples["weight_g"], **rise_options)
    last_s = session.duration_s
    rows = []
    for shift in range(0, int(max_shift_days) + 1):
        end = session.end_time(shift)
        start = end - dt.timedelta(seconds=last_s)
        onsets = [end - dt.timedelta(seconds=last_s - s) for s in rises["onset_s"]]
        used: set[int] = set()
        offsets = []
        for onset in onsets:
            if not events:
                break
            distances = [abs((onset - event).total_seconds()) for event in events]
            best = int(np.argmin(distances))
            if distances[best] <= tolerance_s and best not in used:
                used.add(best)
                offsets.append((onset - events[best]).total_seconds())
        inside = [i for i, event in enumerate(events) if start <= event <= end]
        rows.append(
            {
                "day_shift": shift,
                "start": start,
                "end": end,
                "n_rises": int(len(onsets)),
                "n_matched": int(len(offsets)),
                "n_events_without_rise": int(len([i for i in inside if i not in used])),
                "median_offset_s": (
                    float(np.median(offsets)) if offsets else float("nan")
                ),
                "max_abs_offset_s": (
                    float(np.max(np.abs(offsets))) if offsets else float("nan")
                ),
            }
        )
    frame = pd.DataFrame(rows)
    frame["_abs"] = frame["median_offset_s"].abs().fillna(np.inf)
    frame = frame.sort_values(
        ["n_matched", "n_events_without_rise", "_abs", "day_shift"],
        ascending=[False, True, True, True],
        kind="mergesort",
    )
    return frame.drop(columns="_abs").reset_index(drop=True)


def stitch_sessions(
    anchored: Sequence[AnchoredSession],
    *,
    overlap_tolerance_s: float = 0.5,
    gap_threshold_s: float = 10.0,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Merge anchored sessions into one time series.

    Sessions are taken in the given order of priority. Samples of a later
    session that fall inside the time span of an earlier one are dropped, so
    duplicated exports (for example a sheet copied into a collected workbook)
    are counted once. Returns ``(samples, sessions, gaps)``:

    * ``samples``: ``datetime``, ``weight_g``, ``session`` (label), ``num``;
    * ``sessions``: one row per input with span, kept and dropped samples;
    * ``gaps``: intervals longer than ``gap_threshold_s`` without samples,
      with the weight on both sides, so that level changes across gaps can be
      reviewed.
    """
    tolerance = pd.Timedelta(seconds=overlap_tolerance_s)
    covered: list[tuple[pd.Timestamp, pd.Timestamp]] = []
    kept_frames = []
    session_rows = []
    for item in anchored:
        frame = item.frame
        times = frame["datetime"]
        dropped = np.zeros(len(frame), dtype=bool)
        for start, end in covered:
            dropped |= (
                (times >= start - tolerance) & (times <= end + tolerance)
            ).to_numpy()
        kept = frame.loc[~dropped].copy()
        kept["session"] = item.label
        kept_frames.append(kept[["datetime", "weight_g", "session", "num"]])
        covered.append((item.start, item.end))
        session_rows.append(
            {
                "session": item.label,
                "source_path": item.session.source_path,
                "sheet": item.session.sheet,
                "export_time": item.session.export_time,
                "day_shift": item.day_shift,
                "day_shift_certain": item.session.day_shift_is_certain(),
                "start": item.start,
                "end": item.end,
                "n_samples": int(len(frame)),
                "n_kept": int((~dropped).sum()),
                "n_dropped_overlap": int(dropped.sum()),
                "first_weight_g": float(frame["weight_g"].iloc[0]),
                "last_weight_g": float(frame["weight_g"].iloc[-1]),
            }
        )
    samples = (
        pd.concat(kept_frames, ignore_index=True)
        .sort_values("datetime", kind="mergesort")
        .reset_index(drop=True)
        if kept_frames
        else pd.DataFrame(columns=["datetime", "weight_g", "session", "num"])
    )
    gap_rows = []
    if len(samples) > 1:
        seconds = (
            (samples["datetime"] - samples["datetime"].iloc[0])
            .dt.total_seconds()
            .to_numpy()
        )
        for index in np.flatnonzero(np.diff(seconds) > gap_threshold_s):
            before, after = samples.iloc[index], samples.iloc[index + 1]
            gap_rows.append(
                {
                    "start": before["datetime"],
                    "end": after["datetime"],
                    "duration_s": float(seconds[index + 1] - seconds[index]),
                    "weight_before_g": float(before["weight_g"]),
                    "weight_after_g": float(after["weight_g"]),
                    "delta_g": float(after["weight_g"] - before["weight_g"]),
                    "session_before": before["session"],
                    "session_after": after["session"],
                }
            )
    gaps = pd.DataFrame(
        gap_rows,
        columns=[
            "start",
            "end",
            "duration_s",
            "weight_before_g",
            "weight_after_g",
            "delta_g",
            "session_before",
            "session_after",
        ],
    )
    return samples, pd.DataFrame(session_rows), gaps


def apply_offsets(
    samples: pd.DataFrame, offsets: Iterable[Mapping[str, Any]]
) -> pd.DataFrame:
    """Add documented offsets to ``weight_g`` and keep the raw value.

    Each offset is a mapping with ``start``, optional ``end`` (exclusive) and
    ``offset_g``. Offsets are meant for reviewed events such as a swapped
    container, never for automatic smoothing.
    """
    result = samples.copy()
    if "weight_raw_g" not in result:
        result["weight_raw_g"] = result["weight_g"]
    correction = np.zeros(len(result))
    for offset in offsets:
        start = pd.Timestamp(offset["start"])
        mask = (result["datetime"] >= start).to_numpy()
        if offset.get("end") is not None:
            mask &= (result["datetime"] < pd.Timestamp(offset["end"])).to_numpy()
        correction[mask] += float(offset["offset_g"])
    result["offset_g"] = correction
    result["weight_g"] = result["weight_raw_g"] + correction
    return result


def resample_mean(samples: pd.DataFrame, period_s: float = 10.0) -> pd.DataFrame:
    """Average the samples in fixed bins; time stamps are bin centres."""
    period = pd.Timedelta(seconds=period_s)
    bins = samples["datetime"].dt.floor(period)
    grouped = samples.groupby(bins, sort=True)
    columns = [c for c in samples.columns if c.endswith("_g")]
    result = grouped[columns].mean()
    result.insert(0, "n_samples", grouped.size())
    result = result.reset_index(names="bin_start")
    result.insert(0, "datetime", result["bin_start"] + period / 2)
    return result.drop(columns="bin_start")


def _main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = parser.add_subparsers(dest="command", required=True)
    inspect = sub.add_parser(
        "inspect", help="Summarise the sessions in iWeight exports"
    )
    inspect.add_argument("paths", nargs="+", type=Path)
    inspect.add_argument(
        "--json", action="store_true", help="One JSON object per session"
    )
    args = parser.parse_args(argv)
    for path in args.paths:
        try:
            sessions = read_kinetics_workbook(path)
        except Exception as exc:  # pragma: no cover - command-line reporting
            print(f"{path}: ERROR {exc}", file=sys.stderr)
            continue
        if not sessions:
            print(f"{path}: no iWeight KINETICS sheets")
        for session in sessions:
            info = session.describe()
            if args.json:
                print(json.dumps(info, ensure_ascii=False))
            else:
                print(
                    f"{Path(path).name} [{session.sheet}] export {info['export_time']} "
                    f"n={info['n_samples']} {info['start_time']} -> {info['end_time']} "
                    f"gap={info['export_gap_s']:.0f}s max={info['max_time_s']} "
                    f"reached_max={info['reached_max_time']} "
                    f"certain={info['day_shift_certain']} "
                    f"ignored={info['ignored_rows']} "
                    f"w={info['first_weight_g']}->{info['last_weight_g']}"
                )
    return 0


if __name__ == "__main__":
    raise SystemExit(_main())
