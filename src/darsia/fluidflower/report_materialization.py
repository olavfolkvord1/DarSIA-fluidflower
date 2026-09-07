"""Materialize immutable FluidFlower reporting plans into review packages."""

from __future__ import annotations

import json
import mimetypes
import os
import re
import shutil
import time
import zipfile
from collections.abc import Mapping
from copy import deepcopy
from html import escape
from pathlib import Path
from typing import Any
from urllib.parse import quote

from darsia.fluidflower.execution import (
    AdapterDescriptor,
    TaskExecutionContext,
    load_task_plan,
    verify_file_reference,
)
from darsia.fluidflower.project import FluidFlowerProject, sha256_file, utc_now
from darsia.fluidflower.reporting import (
    REPORTING_REVISION_FORMAT,
    REPORTING_TASK_CONTRACT_FORMAT,
)

REPORT_PACKAGE_FORMAT = "darsia-fluidflower-report-package-v1"
REPORT_PROVENANCE_FORMAT = "darsia-fluidflower-report-provenance-v1"

SECTION_TITLES = {
    "executive_summary": "Executive summary",
    "experimental_setup": "Experimental setup",
    "data_integrity": "Data integrity",
    "geometry_facies": "Geometry and facies",
    "image_correction": "Image correction",
    "temporal_alignment": "Temporal alignment",
    "phase_interpretation": "Phase interpretation",
    "calibration": "Calibration",
    "mass_analysis": "Mass analysis",
    "wasserstein_analysis": "Wasserstein analysis",
    "post_analysis": "Post-analysis",
    "uncertainty_limitations": "Uncertainty and limitations",
    "reproducibility_provenance": "Reproducibility and provenance",
    "approvals": "Approvals",
    "appendices": "Appendices",
}


def _write_json(path: Path, payload: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{os.getpid()}.{time.time_ns()}.tmp")
    try:
        temporary.write_text(
            json.dumps(dict(payload), ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        os.replace(temporary, path)
    finally:
        try:
            temporary.unlink()
        except (FileNotFoundError, OSError):
            pass


def _safe_name(value: str, fallback: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", str(value).strip()).strip("-.")
    return cleaned[:120] or fallback


def _verify_file(path_text: str, expected_sha256: str, label: str) -> Path:
    return verify_file_reference(path_text, expected_sha256, label=label)


def _reference_rows(contract: Mapping[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    post = dict(contract.get("post_analysis_result") or {})
    if post.get("path"):
        rows.append(
            {
                "reference_id": "post-analysis-result-manifest",
                "kind": "post_analysis_result_manifest",
                "title": "Post-analysis result manifest",
                "path": str(post.get("path") or ""),
                "sha256": str(post.get("sha256") or ""),
                "media_type": "application/json",
            }
        )
    for artifact in post.get("artifacts") or []:
        if not isinstance(artifact, Mapping):
            continue
        rows.append(
            {
                "reference_id": str(artifact.get("artifact_id") or "post-artifact"),
                "kind": "post_analysis_artifact",
                "title": str(artifact.get("product") or "Post-analysis artifact"),
                "path": str(artifact.get("path") or ""),
                "sha256": str(artifact.get("sha256") or ""),
                "media_type": str(artifact.get("media_type") or ""),
                "product": str(artifact.get("product") or ""),
                "run_ids": list(artifact.get("run_ids") or []),
            }
        )
    for evidence in contract.get("evidence_files") or []:
        if not isinstance(evidence, Mapping):
            continue
        rows.append(
            {
                "reference_id": str(evidence.get("evidence_id") or "evidence"),
                "kind": str(evidence.get("kind") or "other"),
                "title": str(evidence.get("title") or "Evidence"),
                "path": str(evidence.get("path") or ""),
                "sha256": str(evidence.get("sha256") or ""),
                "media_type": str(evidence.get("media_type") or ""),
            }
        )
    return rows


def _make_thumbnail(source: Path, destination: Path, media_type: str) -> bool:
    try:
        from PIL import Image

        if media_type.startswith("image/"):
            with Image.open(source) as image:
                image.thumbnail((960, 540))
                destination.parent.mkdir(parents=True, exist_ok=True)
                image.convert("RGB").save(destination, "JPEG", quality=88)
            return True
        if media_type.startswith("video/"):
            import cv2

            capture = cv2.VideoCapture(str(source))
            try:
                ok, frame = capture.read()
            finally:
                capture.release()
            if not ok:
                return False
            frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
            image = Image.fromarray(frame)
            image.thumbnail((960, 540))
            destination.parent.mkdir(parents=True, exist_ok=True)
            image.save(destination, "JPEG", quality=88)
            return True
    except Exception:  # noqa: BLE001
        return False
    return False


def _stage_references(
    staging: Path,
    contract: Mapping[str, Any],
) -> list[dict[str, Any]]:
    exports = dict(contract.get("exports") or {})
    maximum_embedded = int(exports.get("maximum_embedded_media_bytes") or 0)
    rows: list[dict[str, Any]] = []
    seen: set[str] = set()
    for index, raw in enumerate(_reference_rows(contract), start=1):
        reference_id = _safe_name(str(raw["reference_id"]), f"reference-{index}")
        if reference_id in seen:
            reference_id = f"{reference_id}-{index}"
        seen.add(reference_id)
        source = _verify_file(
            str(raw.get("path") or ""),
            str(raw.get("sha256") or ""),
            str(raw.get("title") or reference_id),
        )
        media_type = str(raw.get("media_type") or "").strip()
        if not media_type:
            media_type = (
                mimetypes.guess_type(source.name)[0] or "application/octet-stream"
            )
        row = deepcopy(dict(raw))
        row.update(
            {
                "reference_id": reference_id,
                "path": str(source),
                "sha256": sha256_file(source),
                "size_bytes": source.stat().st_size,
                "media_type": media_type,
                "packaged_path": "",
                "thumbnail_path": "",
            }
        )
        if source.stat().st_size <= maximum_embedded:
            packaged_name = (
                f"{reference_id}-{row['sha256'][:12]}-"
                f"{_safe_name(source.name, 'file')}"
            )
            destination = staging / "references" / packaged_name
            destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, destination)
            if sha256_file(destination) != row["sha256"]:
                raise RuntimeError(f"Copied reference failed verification: {source}")
            row["packaged_path"] = destination.relative_to(staging).as_posix()
        else:
            thumbnail = staging / "thumbnails" / f"{reference_id}.jpg"
            if _make_thumbnail(source, thumbnail, media_type):
                row["thumbnail_path"] = thumbnail.relative_to(staging).as_posix()
        rows.append(row)
    return rows


def _module_quality(contract: Mapping[str, Any], module: str) -> dict[str, Any]:
    for row in contract.get("quality_snapshot") or []:
        if isinstance(row, Mapping) and row.get("module") == module:
            return dict(row)
    return {"module": module, "status": "not_recorded", "reason": ""}


def _section_payloads(
    report: Mapping[str, Any],
) -> list[dict[str, Any]]:
    contract = dict(report["task_contract"])
    project = dict(report["project"])
    references = list(report["references"])
    scientific = dict(contract.get("scientific_contract") or {})
    payloads: list[dict[str, Any]] = []
    for section in contract.get("sections") or []:
        title = SECTION_TITLES.get(str(section), str(section).replace("_", " ").title())
        paragraphs: list[str] = []
        tables: list[dict[str, Any]] = []
        if section == "executive_summary":
            paragraphs.extend(
                [
                    (
                        "This package assembles previously computed and hash-verified "
                        "FluidFlower results. It does not alter phase maps, "
                        "concentrations, masses, or transport distances."
                    ),
                    (
                        "Detected raw mass remains the authoritative image-analysis "
                        "result; injected mass is retained as a reference and is not "
                        "imposed as a conservation constraint."
                    ),
                ]
            )
        elif section == "experimental_setup":
            rig = dict(project.get("rig") or {})
            chemistry = dict(project.get("chemistry") or {})
            tables.append(
                {
                    "headers": ["Field", "Recorded value"],
                    "rows": [
                        ["Series", str(project.get("name") or "")],
                        ["Series ID", str(project.get("series_id") or "")],
                        [
                            "Rig type",
                            str(rig.get("type_label") or rig.get("type") or ""),
                        ],
                        [
                            "Indicator system",
                            str(chemistry.get("indicator_system") or ""),
                        ],
                        ["Run count", str(len(project.get("runs") or []))],
                    ],
                }
            )
        elif section in {
            "data_integrity",
            "geometry_facies",
            "image_correction",
            "temporal_alignment",
            "phase_interpretation",
            "calibration",
            "mass_analysis",
            "post_analysis",
        }:
            module = {
                "geometry_facies": "geometry_facies",
                "temporal_alignment": "temporal_sync",
            }.get(str(section), str(section))
            quality = _module_quality(contract, module)
            paragraphs.append(
                f"Frozen quality status: {quality.get('status', 'not_recorded')}. "
                f"{quality.get('reason', '')}".strip()
            )
            if section == "phase_interpretation":
                paragraphs.append(
                    "Phase labels describe the camera-visible 2D slice. Conversion to "
                    "physical volume uses the approved geometry, depth, and porosity."
                )
            if section == "mass_analysis":
                paragraphs.append(
                    "Raw detected mass is reported without forcing agreement with "
                    "injected mass. Phase and mass uncertainty remain separate."
                )
            if section == "post_analysis":
                post_rows = [
                    [
                        str(row.get("title") or ""),
                        ", ".join(str(value) for value in row.get("run_ids") or [])
                        or "All",
                        str(row.get("media_type") or ""),
                        str(row.get("sha256") or "")[:16],
                    ]
                    for row in references
                    if row.get("kind") == "post_analysis_artifact"
                ]
                tables.append(
                    {
                        "headers": ["Product", "Runs", "Media type", "SHA-256"],
                        "rows": post_rows,
                    }
                )
        elif section == "wasserstein_analysis":
            evidence = [
                row for row in references if row.get("kind") == "wasserstein_result"
            ]
            paragraphs.append(
                "Wasserstein results are included as verified evidence and retain "
                "their own resolution and normalization provenance."
            )
            tables.append(
                {
                    "headers": ["Evidence", "Path", "SHA-256"],
                    "rows": [
                        [
                            str(row.get("title") or ""),
                            str(row.get("path") or ""),
                            str(row.get("sha256") or "")[:16],
                        ]
                        for row in evidence
                    ],
                }
            )
        elif section == "uncertainty_limitations":
            tables.append(
                {
                    "headers": [
                        "Category",
                        "Statement",
                        "Effect",
                        "Status",
                        "Reviewer",
                    ],
                    "rows": [
                        [
                            str(row.get("category") or ""),
                            str(row.get("statement") or ""),
                            str(row.get("effect_on_results") or ""),
                            str(row.get("review_status") or ""),
                            str(row.get("reviewed_by") or ""),
                        ]
                        for row in contract.get("uncertainty_register") or []
                        if isinstance(row, Mapping)
                    ],
                }
            )
        elif section == "reproducibility_provenance":
            paragraphs.append(
                "Every task, plan, source manifest, evidence file, and generated "
                "artifact is identified by path and SHA-256 digest in the package "
                "manifests."
            )
            tables.append(
                {
                    "headers": ["Contract", "Value"],
                    "rows": [
                        ["Plan", str(report.get("plan_id") or "")],
                        ["Task", str(report.get("task_id") or "")],
                        ["Backend", str(report.get("backend") or "")],
                        [
                            "Visible scope",
                            str(scientific.get("visible_phase_scope") or ""),
                        ],
                        [
                            "Physical volume",
                            str(scientific.get("physical_volume_source") or ""),
                        ],
                    ],
                }
            )
        elif section == "approvals":
            tables.append(
                {
                    "headers": ["Role", "Name", "Decision", "Date", "Note"],
                    "rows": [
                        [
                            str(row.get("role") or ""),
                            str(row.get("name") or ""),
                            str(row.get("decision") or ""),
                            str(row.get("date") or ""),
                            str(row.get("note") or ""),
                        ]
                        for row in contract.get("signoffs") or []
                        if isinstance(row, Mapping)
                    ],
                }
            )
        elif section == "appendices":
            tables.append(
                {
                    "headers": ["Evidence", "Type", "Stored as", "SHA-256"],
                    "rows": [
                        [
                            str(row.get("title") or ""),
                            str(row.get("kind") or ""),
                            str(row.get("packaged_path") or row.get("path") or ""),
                            str(row.get("sha256") or "")[:16],
                        ]
                        for row in references
                    ],
                }
            )
        payloads.append(
            {
                "key": str(section),
                "title": title,
                "paragraphs": paragraphs,
                "tables": tables,
            }
        )
    return payloads


def _html_link(row: Mapping[str, Any]) -> str:
    packaged = str(row.get("packaged_path") or "")
    if packaged:
        return quote(packaged)
    path = Path(str(row.get("path") or "")).expanduser().resolve()
    try:
        return path.as_uri()
    except ValueError:
        return ""


def _render_html(report: Mapping[str, Any], destination: Path) -> None:
    identity = dict(report["task_contract"].get("identity") or {})
    section_markup: list[str] = []
    for section in report["section_payloads"]:
        body = "".join(f"<p>{escape(value)}</p>" for value in section["paragraphs"])
        for table in section["tables"]:
            headers = "".join(f"<th>{escape(value)}</th>" for value in table["headers"])
            rows = "".join(
                "<tr>" + "".join(f"<td>{escape(value)}</td>" for value in row) + "</tr>"
                for row in table["rows"]
            )
            body += (
                "<div class='table-wrap'><table><thead><tr>"
                f"{headers}</tr></thead><tbody>{rows}</tbody></table></div>"
            )
        section_markup.append(
            f"<section id='{escape(section['key'])}'>"
            f"<h2>{escape(section['title'])}</h2>{body}</section>"
        )
    evidence_items: list[str] = []
    for row in report["references"]:
        href = escape(_html_link(row), quote=True)
        title = escape(str(row.get("title") or "Evidence"))
        digest = escape(str(row.get("sha256") or "")[:16])
        evidence_items.append(
            f"<li><a href='{href}'>{title}</a> <code>{digest}</code></li>"
        )
    evidence_links = "".join(evidence_items)
    watermark = "<div class='watermark'>DRAFT</div>" if report["draft"] else ""
    destination.write_text(
        "<!doctype html><html lang='en'><head><meta charset='utf-8'>"
        "<meta name='viewport' content='width=device-width,initial-scale=1'>"
        f"<title>{escape(str(identity.get('title') or 'FluidFlower report'))}</title>"
        "<style>body{font-family:Arial,sans-serif;margin:0;color:#18202a;background:#fff}"
        "main{max-width:1100px;margin:auto;padding:48px 56px}"
        "h1{font-size:34px;margin:0 0 8px}"
        "h2{font-size:22px;border-bottom:2px solid #1f6f78;"
        "padding-bottom:7px;margin-top:38px}"
        "p{line-height:1.55;max-width:85ch}.meta{color:#52606d;margin-bottom:28px}"
        ".table-wrap{overflow:auto}table{border-collapse:collapse;"
        "width:100%;font-size:13px}"
        "th,td{border:1px solid #cbd3da;padding:8px;text-align:left;vertical-align:top}"
        "th{background:#edf3f3}code{font-size:12px}"
        ".watermark{position:fixed;right:24px;top:18px;"
        "font:700 18px Arial;color:#a13b3b;letter-spacing:2px}"
        "@media(max-width:700px){main{padding:24px}}"
        "</style></head><body>"
        f"{watermark}<main><h1>"
        f"{escape(str(identity.get('title') or 'FluidFlower report'))}</h1>"
        f"<div class='meta'>{escape(str(identity.get('subtitle') or ''))}<br>"
        f"Version {escape(str(identity.get('report_version') or ''))} | "
        f"Generated {escape(str(report.get('generated_utc') or ''))}</div>"
        + "".join(section_markup)
        + "<section><h2>Verified evidence links</h2><ul>"
        + evidence_links
        + "</ul></section>"
        "</main></body></html>",
        encoding="utf-8",
    )


def _require_dependency(module: str, package: str) -> Any:
    try:
        return __import__(module)
    except ImportError as exc:
        raise RuntimeError(
            f"The selected report export requires {package}. Install the "
            "'fluidflower-app' optional dependencies."
        ) from exc


def _render_pdf(report: Mapping[str, Any], destination: Path) -> None:
    _require_dependency("reportlab", "reportlab")
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    identity = dict(report["task_contract"].get("identity") or {})
    styles = getSampleStyleSheet()
    version_line = (
        f"Version {identity.get('report_version', '')} | "
        f"Generated {report.get('generated_utc', '')}"
    )
    story: list[Any] = [
        Paragraph(
            escape(str(identity.get("title") or "FluidFlower report")), styles["Title"]
        ),
        Paragraph(escape(str(identity.get("subtitle") or "")), styles["Heading2"]),
        Paragraph(
            escape(version_line),
            styles["Normal"],
        ),
        Spacer(1, 8 * mm),
    ]
    for index, section in enumerate(report["section_payloads"]):
        if index:
            story.append(PageBreak())
        story.append(Paragraph(escape(section["title"]), styles["Heading1"]))
        for value in section["paragraphs"]:
            story.append(Paragraph(escape(value), styles["BodyText"]))
            story.append(Spacer(1, 2 * mm))
        for table in section["tables"]:
            data = [table["headers"], *table["rows"]]
            wrapped = [
                [Paragraph(escape(str(cell)), styles["BodyText"]) for cell in row]
                for row in data
            ]
            pdf_table = Table(wrapped, repeatRows=1)
            pdf_table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dfecec")),
                        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#aab5bd")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("FONTSIZE", (0, 0), (-1, -1), 7),
                    ]
                )
            )
            story.append(pdf_table)
            story.append(Spacer(1, 4 * mm))
    document = SimpleDocTemplate(
        str(destination),
        pagesize=landscape(A4),
        rightMargin=14 * mm,
        leftMargin=14 * mm,
    )
    document.build(story)


def _render_docx(report: Mapping[str, Any], destination: Path) -> None:
    _require_dependency("docx", "python-docx")
    from docx import Document

    identity = dict(report["task_contract"].get("identity") or {})
    document = Document()
    document.add_heading(str(identity.get("title") or "FluidFlower report"), 0)
    document.add_paragraph(str(identity.get("subtitle") or ""))
    version_line = (
        f"Version {identity.get('report_version', '')} | "
        f"Generated {report.get('generated_utc', '')}"
    )
    document.add_paragraph(version_line)
    for section in report["section_payloads"]:
        document.add_heading(section["title"], level=1)
        for value in section["paragraphs"]:
            document.add_paragraph(value)
        for table in section["tables"]:
            word_table = document.add_table(rows=1, cols=max(1, len(table["headers"])))
            word_table.style = "Table Grid"
            for index, value in enumerate(table["headers"]):
                word_table.rows[0].cells[index].text = value
            for row in table["rows"]:
                cells = word_table.add_row().cells
                for index, value in enumerate(row):
                    cells[index].text = value
    document.save(destination)


def _render_pptx(report: Mapping[str, Any], destination: Path) -> None:
    _require_dependency("pptx", "python-pptx")
    from pptx import Presentation
    from pptx.util import Inches, Pt

    identity = dict(report["task_contract"].get("identity") or {})
    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = str(identity.get("title") or "FluidFlower report")
    title_slide.placeholders[1].text = (
        f"{identity.get('subtitle', '')}\nVersion {identity.get('report_version', '')}"
    )
    for section in report["section_payloads"]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = section["title"]
        text_box = slide.shapes.add_textbox(
            Inches(0.7), Inches(1.35), Inches(11.8), Inches(5.5)
        )
        frame = text_box.text_frame
        frame.word_wrap = True
        lines = list(section["paragraphs"])
        for table in section["tables"]:
            lines.append(" | ".join(table["headers"]))
            lines.extend(" | ".join(row) for row in table["rows"][:8])
            if len(table["rows"]) > 8:
                remaining = len(table["rows"]) - 8
                lines.append(f"... {remaining} additional rows in the report package")
        for index, value in enumerate(
            lines or ["See the report package for verified artifacts."]
        ):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = value
            paragraph.font.size = Pt(17 if index < 2 else 13)
            paragraph.space_after = Pt(8)
    presentation.save(destination)


def _file_manifest(
    root: Path, *, exclude: set[Path] | None = None
) -> list[dict[str, Any]]:
    excluded = {path.resolve() for path in (exclude or set())}
    rows: list[dict[str, Any]] = []
    for path in sorted(
        candidate for candidate in root.rglob("*") if candidate.is_file()
    ):
        if path.resolve() in excluded:
            continue
        rows.append(
            {
                "path": path.relative_to(root).as_posix(),
                "size_bytes": path.stat().st_size,
                "sha256": sha256_file(path),
            }
        )
    return rows


def _write_provenance_zip(staging: Path, destination: Path) -> None:
    with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for row in _file_manifest(staging, exclude={destination}):
            archive.write(staging / row["path"], row["path"])


class ReportPackageAdapter:
    """Render a verified reporting plan without changing scientific results."""

    descriptor = AdapterDescriptor(
        key="report-package-v1",
        version="1.0.0",
        title="Scientific report package",
        modules=("reporting",),
        backends=("cpu-float64",),
        description=(
            "Builds HTML, PDF, DOCX, PPTX, JSON, and provenance exports from a "
            "saved reporting plan and hash-verified post-analysis artifacts."
        ),
    )

    def execute(
        self,
        task: Mapping[str, Any],
        context: TaskExecutionContext,
    ) -> Mapping[str, Any]:
        if task.get("module") != "reporting":
            raise ValueError("The report package adapter only accepts reporting tasks")
        if (task.get("backend") or {}).get("name") != "cpu-float64":
            raise ValueError("Report packages require the cpu-float64 backend contract")
        plan, plan_path, plan_file_sha256 = load_task_plan(task)
        if plan.get("format") != REPORTING_REVISION_FORMAT:
            raise ValueError("The task does not reference a reporting revision")
        contract = dict(plan.get("task_contract") or {})
        if contract.get("format") != REPORTING_TASK_CONTRACT_FORMAT:
            raise ValueError("The reporting task contract is unsupported")
        if not bool((plan.get("validation") or {}).get("generation_ready")):
            raise RuntimeError("The saved reporting plan is not generation-ready")

        project = FluidFlowerProject.load(Path(str(task["project"]["root"])))
        if project.project_id != str(task["project"].get("project_id") or ""):
            raise RuntimeError("Task and project identities do not match")
        source_manifest_sha256 = str(
            (task.get("project") or {}).get("source_manifest_sha256") or ""
        )
        if source_manifest_sha256:
            _verify_file(
                str(project.source_manifest_path),
                source_manifest_sha256,
                "Project source manifest",
            )

        context.checkpoint({"stage": "verifying_references"})
        references = _stage_references(context.staging, contract)
        context.checkpoint(
            {"stage": "references_verified", "reference_count": len(references)}
        )

        identity = dict(contract.get("identity") or {})
        formats = [
            str(value) for value in (contract.get("exports") or {}).get("formats") or []
        ]
        report: dict[str, Any] = {
            "schema": 1,
            "format": REPORT_PACKAGE_FORMAT,
            "package_id": str(task["task_id"]),
            "generated_utc": utc_now(),
            "generated_by": str(task.get("created_by") or ""),
            "task_id": str(task["task_id"]),
            "task_sha256": str(task.get("content_sha256") or ""),
            "plan_id": str(plan.get("plan_id") or ""),
            "plan_path": str(plan_path),
            "plan_file_sha256": plan_file_sha256,
            "plan_content_sha256": str(plan.get("content_sha256") or ""),
            "backend": str((task.get("backend") or {}).get("name") or ""),
            "adapter": deepcopy(dict(task.get("adapter") or {})),
            "project": deepcopy(dict(project.manifest)),
            "project_manifest_sha256_at_execution": sha256_file(project.manifest_path),
            "source_manifest_sha256": (
                sha256_file(project.source_manifest_path)
                if project.source_manifest_path.is_file()
                else ""
            ),
            "task_contract": deepcopy(contract),
            "references": references,
            "draft": bool((contract.get("exports") or {}).get("draft_watermark", True)),
            "publication_ready_at_plan_save": bool(
                (plan.get("validation") or {}).get("publication_ready")
            ),
            "scientific_results_created": False,
            "selected_exports": formats,
        }
        report["section_payloads"] = _section_payloads(report)

        _write_json(context.staging / "report.json", report)
        _write_json(context.staging / "task.json", task)
        _write_json(context.staging / "plan.json", plan)
        _write_json(context.staging / "hardware.json", context.inventory)
        shutil.copy2(project.manifest_path, context.staging / "series.json")
        if project.source_manifest_path.is_file():
            shutil.copy2(
                project.source_manifest_path,
                context.staging / "source_manifest.json",
            )

        renderers = {
            "html": ("index.html", _render_html),
            "pdf": ("report.pdf", _render_pdf),
            "docx": ("report.docx", _render_docx),
            "pptx": ("report.pptx", _render_pptx),
        }
        generated: list[str] = ["report.json"]
        for report_format, (filename, renderer) in renderers.items():
            if report_format not in formats:
                continue
            if context.cancellation_requested():
                raise RuntimeError("Report generation was cancelled")
            context.checkpoint({"stage": f"rendering_{report_format}"})
            renderer(report, context.staging / filename)
            generated.append(filename)

        provenance = {
            "schema": 1,
            "format": REPORT_PROVENANCE_FORMAT,
            "task_id": str(task["task_id"]),
            "plan_id": str(plan.get("plan_id") or ""),
            "generated_utc": report["generated_utc"],
            "references": references,
            "files": _file_manifest(context.staging),
        }
        _write_json(context.staging / "provenance_manifest.json", provenance)
        if "provenance_bundle_zip" in formats:
            _write_provenance_zip(
                context.staging, context.staging / "provenance_bundle.zip"
            )
            generated.append("provenance_bundle.zip")

        package_manifest = {
            "schema": 1,
            "format": REPORT_PACKAGE_FORMAT,
            "package_id": str(task["task_id"]),
            "title": str(identity.get("title") or "FluidFlower report"),
            "draft": report["draft"],
            "publication_ready_at_plan_save": report["publication_ready_at_plan_save"],
            "generated_exports": generated,
            "files": _file_manifest(context.staging),
        }
        _write_json(context.staging / "package_manifest.json", package_manifest)
        context.checkpoint(
            {
                "stage": "report_package_ready",
                "generated_exports": generated,
                "file_count": len(package_manifest["files"]) + 1,
            }
        )
        return {
            "package_id": str(task["task_id"]),
            "title": str(identity.get("title") or "FluidFlower report"),
            "draft": report["draft"],
            "publication_ready_at_plan_save": report["publication_ready_at_plan_save"],
            "generated_exports": generated,
            "reference_count": len(references),
            "scientific_results_created": False,
        }
